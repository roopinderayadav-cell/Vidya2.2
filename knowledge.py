"""
knowledge.py - Vidya's knowledge base and answer engine.

  * stores documents per region (SQLite + files on disk), reads PDF/Word/PPT/Excel,
    tables, and (with AI) screenshots/images/scans inside documents
  * Layer 2 search: word matching + Gemini semantic search across ALL documents
  * runs the 3 layers for every question and decides: answered / not found / conflict
  * SELF-LEARNING:
      - 👍 answers are saved automatically and reused for similar questions
      - 👍 / 👎 on a passage raises / lowers it for similar questions in future
      - 👎 on an auto-learned answer switches that answer off
      - after a new upload, open knowledge gaps are re-checked automatically
  * knowledge gaps, analytics data and the Excel export
"""
import io
import json
import re
import sqlite3
import time
import uuid
from datetime import datetime, timedelta

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

import brain
import config
import llm
import privacy
import rerank
import farerule

# ===========================================================================
# 1. DATABASE
# ===========================================================================

def _conn():
    config.DATA_DIR.mkdir(exist_ok=True)
    config.UPLOAD_DIR.mkdir(exist_ok=True)
    con = sqlite3.connect(config.DB_PATH, check_same_thread=False, timeout=30)
    con.row_factory = sqlite3.Row
    return con


def _add_cols(con, table, cols):
    have = [r[1] for r in con.execute(f"PRAGMA table_info({table})")]
    for name, kind in cols:
        if name not in have:
            con.execute(f"ALTER TABLE {table} ADD COLUMN {name} {kind}")


def init_db():
    with _conn() as con:
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                region TEXT, doc_type TEXT, title TEXT, filename TEXT,
                stored_path TEXT, uploaded_by TEXT, uploaded_at TEXT,
                version TEXT, chunks INTEGER DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                doc_id INTEGER, region TEXT, location TEXT, text TEXT
            );
            CREATE TABLE IF NOT EXISTS questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                region TEXT, asked_by TEXT, question TEXT, answer TEXT,
                sources TEXT, score REAL, feedback INTEGER, comment TEXT,
                status TEXT, asked_at TEXT
            );
            CREATE TABLE IF NOT EXISTS learned (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                region TEXT, question TEXT, answer TEXT,
                added_by TEXT, added_at TEXT, from_question INTEGER
            );
            CREATE TABLE IF NOT EXISTS feedback_pairs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                region TEXT, question TEXT, source_key TEXT, value INTEGER, at TEXT
            );
            CREATE TABLE IF NOT EXISTS announcements (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT, body TEXT, level TEXT, region TEXT,
                author TEXT, created_at TEXT, active INTEGER DEFAULT 1
            );
            CREATE TABLE IF NOT EXISTS app_settings (key TEXT PRIMARY KEY, value TEXT);
            CREATE TABLE IF NOT EXISTS sme_chats (
                id INTEGER PRIMARY KEY AUTOINCREMENT, region TEXT, sub_lob TEXT, agent TEXT, sme TEXT,
                status TEXT, started_at TEXT, picked_at TEXT, closed_at TEXT, closed_by TEXT,
                from_question INTEGER, saved INTEGER DEFAULT 0, saved_count INTEGER DEFAULT 0);
            CREATE TABLE IF NOT EXISTS sme_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT, chat_id INTEGER, sender TEXT, role TEXT,
                text TEXT, at TEXT);
            CREATE TABLE IF NOT EXISTS sme_presence (
                user TEXT PRIMARY KEY, region TEXT, available INTEGER, last_seen REAL);
            CREATE TABLE IF NOT EXISTS doc_images (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                doc_id INTEGER, location TEXT, n INTEGER, filename TEXT
            );
            """
        )
        # upgrades from Compass v1 / v2 - existing data is kept
        _add_cols(con, "documents", [("status", "TEXT DEFAULT 'Active'"), ("owner", "TEXT"),
                                     ("keywords", "TEXT"), ("lob", "TEXT DEFAULT ''"),
                                     ("sub_lob", "TEXT DEFAULT ''"), ("images_done", "INTEGER DEFAULT 0")])
        _add_cols(con, "chunks", [("embedding", "BLOB")])
        _add_cols(con, "questions", [("image_path", "TEXT"), ("meta", "TEXT"), ("topic", "TEXT"),
                                     ("result", "TEXT"), ("confidence", "TEXT"),
                                     ("gap_status", "TEXT"), ("gap_owner", "TEXT"),
                                     ("gap_note", "TEXT"), ("lob", "TEXT"), ("sub_lob", "TEXT"),
                                     ("confidence_pct", "INTEGER")])
        _add_cols(con, "questions", [("sme_reason", "TEXT"), ("sme_by", "TEXT"), ("sme_at", "TEXT")])
        _add_cols(con, "learned", [("variants", "TEXT DEFAULT ''")])   # how agents actually asked it
        _add_cols(con, "sme_messages", [("image_path", "TEXT")])
        _add_cols(con, "announcements", [("ticker", "INTEGER DEFAULT 1"), ("until", "TEXT")])
        _add_cols(con, "documents", [("effective_date", "TEXT")])
        _add_cols(con, "learned", [("source", "TEXT DEFAULT 'verified'"),
                                   ("active", "INTEGER DEFAULT 1"), ("votes", "INTEGER DEFAULT 0"),
                                   ("embedding", "BLOB"), ("lob", "TEXT DEFAULT ''")])


# ---------------------------------------------------------------------------
# Lines of business (editable by admins, stored in the database)
# ---------------------------------------------------------------------------
def get_lobs():
    """{LOB: [sub-LOB, ...]} for every LOB tab (config.REGIONS), in display order."""
    stored = {}
    with _conn() as con:
        row = con.execute("SELECT value FROM app_settings WHERE key='lobs'").fetchone()
    if row:
        try:
            data = json.loads(row["value"])
            if isinstance(data, dict):
                stored = {str(k): [str(x) for x in v] for k, v in data.items()}
        except (ValueError, TypeError):
            stored = {}
    return {r: stored.get(r, list(config.DEFAULT_LOBS.get(r, []))) for r in config.REGIONS}


def save_lobs(lobs):
    """lobs = {LOB: [sub-LOB, ...]} for the LOB tabs in config.REGIONS. Sub-LOB names are trimmed,
    blanks and duplicates dropped."""
    clean = {}
    for name, subs in lobs.items():
        name = str(name or "").strip()
        if name not in config.REGIONS:
            raise ValueError(f"'{name}' is not a LOB tab. LOB tabs are: {', '.join(config.REGIONS)}.")
        seen, out = set(), []
        for sub in subs:
            sub = str(sub or "").strip()
            if sub and sub.lower() not in seen:
                seen.add(sub.lower())
                out.append(sub)
        clean[name] = out
    merged = get_lobs()
    merged.update(clean)
    with _conn() as con:
        con.execute("INSERT INTO app_settings (key, value) VALUES ('lobs', ?) "
                    "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (json.dumps(merged),))
    return merged


def update_document_tags(doc_id, sub_lob="", effective_date=None):
    with _conn() as con:
        con.execute("UPDATE documents SET sub_lob=? WHERE id=?", (sub_lob or "", doc_id))
        if effective_date is not None:
            con.execute("UPDATE documents SET effective_date=? WHERE id=?", (effective_date or None, doc_id))


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _vec_to_blob(v):
    return np.asarray(v, dtype=np.float32).tobytes()


def _blob_to_vec(b):
    return np.frombuffer(b, dtype=np.float32) if b else None


# ===========================================================================
# 2. TEXT EXTRACTION
# ===========================================================================

IMAGE_EXTS = ("png", "jpg", "jpeg")


def extract_text(filename: str, data: bytes):
    """Return (parts, images).
    parts  = list of (location, text), e.g. ("2. Create a New PNR", "...")
    images = list of (location, image_bytes) - screenshots/tables inside the file."""
    ext = filename.rsplit(".", 1)[-1].lower()
    parts, images = [], []

    if ext == "pdf":
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    txt = page.extract_text() or ""
                    parts.append((f"Page {i}", txt))
                    for t_i, table in enumerate(page.extract_tables() or [], 1):
                        rows = [" | ".join((c or "").strip().replace("\n", " ") for c in r) for r in table if r]
                        if len(rows) > 1:
                            parts.append((f"Page {i} - Table {t_i}", "\n".join(rows)))
                    if len(txt.strip()) < 30:          # scanned page: read the picture
                        images.append((f"Page {i} (scanned)", _render_pdf_page(data, i - 1)))
        except ImportError:
            from pypdf import PdfReader
            for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, 1):
                parts.append((f"Page {i}", page.extract_text() or ""))
        try:                                            # pictures placed in the PDF
            from pypdf import PdfReader
            for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, 1):
                for img in page.images:
                    images.append((f"Page {i}", img.data))
        except Exception:
            pass

    elif ext == "docx":
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = docx.Document(io.BytesIO(data))
        list_kind = _docx_list_reader(d)
        section, buf, t_i, counter = "Document", [], 0, {}
        for el in d.element.body.iterchildren():
            tag = el.tag.split("}")[-1]
            if tag == "p":
                p = Paragraph(el, d)
                style = (p.style.name if p.style is not None else "").lower()
                if (style.startswith("heading") or style == "title") and p.text.strip():
                    if buf:
                        parts.append((section, "\n".join(buf)))
                    section, buf, counter = p.text.strip(), [p.text.strip()], {}
                elif p.text.strip():
                    kind, key = list_kind(p)
                    text = p.text.strip()
                    if kind == "decimal":       # Word's automatic 1. 2. 3. is not in the text
                        counter[key] = counter.get(key, 0) + 1
                        if not re.match(r"^\d+[.)]", text):
                            text = f"{counter[key]}. {text}"
                    elif kind == "bullet" and not text.startswith(("-", "•")):
                        text = f"- {text}"
                    buf.append(text)
                for blip in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                    rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    if rid and rid in d.part.related_parts:
                        images.append((section, d.part.related_parts[rid].blob))
            elif tag == "tbl":
                t_i += 1
                table = Table(el, d)
                rows = []
                for r in table.rows:
                    cells = []
                    for c in r.cells:
                        if not cells or c.text.strip() != cells[-1]:   # merged cells repeat
                            cells.append(c.text.strip())
                    rows.append(" | ".join(cells))
                parts.append((f"{section} - Table {t_i}", "\n".join(rows)))
        if buf:
            parts.append((section, "\n".join(buf)))

    elif ext == "pptx":
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(io.BytesIO(data))
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for sh in slide.shapes:
                if sh.has_text_frame:
                    texts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    texts += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    images.append((f"Slide {i}", sh.image.blob))
            parts.append((f"Slide {i}", "\n".join(texts)))

    elif ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
            parts.append((f"Sheet {ws.title}", "\n".join(rows)))

    elif ext in IMAGE_EXTS:                             # a screenshot / table uploaded on its own
        images.append(("Image", data))

    else:  # txt, md, csv - split into sections at heading lines
        parts += _text_sections(data.decode("utf-8", errors="ignore"))

    return [(loc, txt) for loc, txt in parts if txt and txt.strip()], images


def _docx_list_reader(d):
    """Returns a function telling whether a Word paragraph is a numbered or bulleted list item."""
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    fmts = {}
    try:
        numbering = d.part.numbering_part.element
        abstract = {a.get(W + "abstractNumId"): a for a in numbering.iter(W + "abstractNum")}
        for num in numbering.iter(W + "num"):
            ref = num.find(W + "abstractNumId")
            a = abstract.get(ref.get(W + "val")) if ref is not None else None
            if a is None:
                continue
            for lvl in a.iter(W + "lvl"):
                f = lvl.find(W + "numFmt")
                fmts[(num.get(W + "numId"), lvl.get(W + "ilvl"))] = f.get(W + "val") if f is not None else ""
    except Exception:
        pass

    def kind(p):
        ppr = p._p.pPr
        num = ppr.numPr if ppr is not None else None
        style = (p.style.name if p.style is not None else "").lower()
        if num is None or num.numId is None:
            if "list number" in style:
                return "decimal", style
            if "list bullet" in style:
                return "bullet", style
            return None, None
        num_id = str(num.numId.val)
        ilvl = str(num.ilvl.val) if num.ilvl is not None else "0"
        fmt = fmts.get((num_id, ilvl), "decimal")
        return ("bullet" if fmt == "bullet" else "decimal"), (num_id, ilvl)
    return kind


def _is_heading(line, nxt):
    t = line.strip()
    if not t or not nxt.strip():
        return False
    if t.startswith("#"):
        return True
    if t.startswith(("Q:", "A:", "-", "*", "•")) or re.match(r"^\d+[.)]\s", t) or "|" in t:
        return False
    return len(t.split()) <= 7 and not t.endswith((".", ":", ";", ",", "?")) and t[0].isupper()


def _text_sections(text):
    """Plain text / markdown: short title lines (or # headings) start a new section."""
    lines = text.splitlines()
    out, section, buf = [], "Text", []
    for i, line in enumerate(lines):
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        if _is_heading(line, nxt):
            if buf:
                out.append((section, "\n".join(buf)))
            section = line.strip().lstrip("#").strip()
            buf = [section]
        elif line.strip():
            buf.append(line.strip())
    if buf:
        out.append((section, "\n".join(buf)))
    return out


def _render_pdf_page(data, index):
    try:
        import pypdfium2 as pdfium
        page = pdfium.PdfDocument(data)[index]
        buf = io.BytesIO()
        page.render(scale=2).to_pil().save(buf, format="PNG")
        return buf.getvalue()
    except Exception:
        return b""


def read_images(title, images):
    """Turn images inside a document into searchable text (needs the AI key).
    Runs several images at the same time so uploads stay quick."""
    if not images or not llm.vision_on() or not config.READ_DOC_IMAGES:
        return [], len(images)
    from concurrent.futures import ThreadPoolExecutor
    images = [(loc, b) for loc, b in images if b][: config.MAX_IMAGES_PER_DOC]

    def work(item):
        loc, b = item
        try:
            return loc, brain.describe_doc_image(b, title, loc)
        except Exception:
            return loc, ""

    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(work, images))
    out, counter = [], {}
    for loc, text in results:
        if text:
            counter[loc] = counter.get(loc, 0) + 1
            out.append((f"{loc} - image {counter[loc]}", "[From image in document]\n" + text))
    return out, 0


def chunk_text(text: str, size=config.CHUNK_SIZE, overlap=config.CHUNK_OVERLAP):
    """Split long text into overlapping pieces that keep paragraphs together.
    FAQ-style text ("Q: ..." lines) is split so each question+answer is its own piece."""
    text = re.sub(r"[ \t]+", " ", text).strip()
    faq = re.split(r"\n(?=\s*(?:Q\s*[:.\-]|Question\s*\d*\s*[:.]))", text, flags=re.I)
    if len(faq) >= 3:
        out = []
        for piece in faq:
            out += chunk_text(piece, size, overlap) if len(piece) > size else [piece.strip()]
        return [o for o in out if o]
    if len(text) <= size:
        return [text]
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    chunks, current = [], ""
    for p in paragraphs:
        if len(current) + len(p) + 1 <= size:
            current = f"{current}\n{p}".strip()
        else:
            if current:
                chunks.append(current)
            while len(p) > size:          # very long paragraph
                chunks.append(p[:size])
                p = p[size - overlap:]
            current = (current[-overlap:] + "\n" + p).strip() if current else p
    if current:
        chunks.append(current)
    return chunks


# ===========================================================================
# 3. DOCUMENT MANAGEMENT
# ===========================================================================

def add_document(region, doc_type, title, filename, data: bytes, uploaded_by, version="", owner="",
                 sub_lob="", effective_date=""):
    """Returns (doc_id, pieces, info). info = images read/skipped, semantic index, gaps re-checked."""
    parts, images = extract_text(filename, data)
    # Card numbers: a REAL card number refuses the upload (nothing is stored, not even the file);
    # public test numbers (4111 1111 1111 1111 ...) are allowed and masked like any other.
    bad = [f"{loc}: {', '.join(privacy.real_cards(privacy.find_cards(txt)))}" for loc, txt in parts
           if privacy.real_cards(privacy.find_cards(txt))]
    images, image_cards = mask_doc_images(images, with_cards=True)
    bad += [f"{loc} (picture): {', '.join(privacy.real_cards(c))}" for loc, c in image_cards
            if privacy.real_cards(c)]
    if bad:
        raise ValueError("Not uploaded - this file contains payment card numbers, which Vidya must not "
                         "store. Remove or mask them in the file and upload again. Found in: "
                         + "; ".join(bad[:5]))
    parts = [(loc, privacy.mask(txt)) for loc, txt in parts]
    image_parts, skipped = read_images(title or filename, images)
    parts += image_parts
    if not parts:
        if images and not llm.vision_on():
            raise ValueError("This file is only images/scans. Reading images needs an AI that "
                             "reads pictures (Gemini/Claude key, or LOCAL_VISION_MODEL).")
        raise ValueError("No readable text found.")

    stored = config.UPLOAD_DIR / f"{uuid.uuid4().hex[:8]}_{filename}"
    stored.write_bytes(data)

    # template guidance sections are never used for answers
    parts = [(loc, "\n".join(l for l in txt.splitlines() if not l.strip().lower().startswith("replace the example")))
             for loc, txt in parts if not loc.lower().startswith(("how to use this template", "instructions"))]
    # the "Document information" table becomes hidden keywords for the whole document
    keywords = " ".join(txt.replace("|", " ") for loc, txt in parts
                        if loc.lower().startswith("document information"))
    parts = [(loc, txt) for loc, txt in parts if not loc.lower().startswith("document information")]
    parts = [(loc, txt) for loc, txt in parts if txt.strip()]
    pieces = []
    for loc, txt in parts:
        for piece in ([txt] if "Table" in loc or "image" in loc else chunk_text(txt)):
            if len(piece) < 40 and len(parts) > 1:      # skip title-only fragments
                continue
            rest = piece.replace(loc.rsplit(" - ", 1)[-1] if " - Table" in loc else loc, "").strip(" \n-:")
            if len(rest) < 15 and len(parts) > 1:        # piece is only its own heading (e.g. before the first Q:)
                continue
            pieces.append((loc, piece))

    vectors, semantic = None, "off"
    check_embed_space()
    if llm.embeddings_on() and pieces:
        try:
            vectors = llm.embed([f"{title or filename} | {loc}\n{p}" for loc, p in pieces])
            semantic = "on"
        except llm.AIError as e:
            semantic = f"failed: {e}"

    with _conn() as con:
        cur = con.execute(
            "INSERT INTO documents (region, doc_type, title, filename, stored_path, uploaded_by,"
            " uploaded_at, version, status, owner, keywords, lob, sub_lob, effective_date)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (region, doc_type, title or filename, filename, str(stored), uploaded_by, _now(),
             version, "Active", owner, keywords[:2000], region, sub_lob or "", effective_date or None))
        doc_id = cur.lastrowid
        for i, (loc, piece) in enumerate(pieces):
            con.execute("INSERT INTO chunks (doc_id, region, location, text, embedding) VALUES (?,?,?,?,?)",
                        (doc_id, region, loc, piece,
                         _vec_to_blob(vectors[i]) if vectors is not None else None))
        con.execute("UPDATE documents SET chunks=? WHERE id=?", (len(pieces), doc_id))
    stored_images = store_images(doc_id, images)
    rechecked = recheck_gaps(region, doc_title=title or filename)
    return doc_id, len(pieces), {"images_found": len(images), "images_read": len(image_parts),
                                 "images_stored": stored_images,
                                 "images_skipped": skipped, "semantic": semantic,
                                 "gaps_rechecked": rechecked}


def reindex_embeddings(progress=None):
    """Build the semantic index for pieces/answers that don't have one yet (after adding a key)."""
    if not llm.embeddings_on():
        raise llm.AIError("Turn on meaning search first (MEANING_SEARCH, a Gemini key or AI_PROVIDER = local)")
    check_embed_space()
    with _conn() as con:
        rows = con.execute("SELECT c.id, c.location, c.text, d.title FROM chunks c JOIN documents d "
                           "ON d.id=c.doc_id WHERE c.embedding IS NULL").fetchall()
        learned = con.execute("SELECT id, question, answer FROM learned WHERE embedding IS NULL").fetchall()
    done = 0
    for i in range(0, len(rows), 50):
        batch = rows[i:i + 50]
        vecs = llm.embed([f"{r['title']} | {r['location']}\n{r['text']}" for r in batch])
        with _conn() as con:
            for r, v in zip(batch, vecs):
                con.execute("UPDATE chunks SET embedding=? WHERE id=?", (_vec_to_blob(v), r["id"]))
        done += len(batch)
        if progress:
            progress(done / max(1, len(rows)))
    if learned:
        vecs = llm.embed([f"{l['question']}\n{l['answer']}" for l in learned])
        with _conn() as con:
            for l, v in zip(learned, vecs):
                con.execute("UPDATE learned SET embedding=? WHERE id=?", (_vec_to_blob(v), l["id"]))
    return done + len(learned)


_space_checked = {"id": None}


def check_embed_space():
    """Meaning-vectors made by one model can't be compared with another model's.
    When the meaning-search model changes (e.g. Gemini -> local), clear the old vectors
    so Settings -> 'Build semantic index now' rebuilds them with the new model.
    Returns True when old vectors were cleared."""
    current = llm.embedder_id()
    if not current or _space_checked["id"] == current:
        return False
    cleared = False
    with _conn() as con:
        row = con.execute("SELECT value FROM app_settings WHERE key='embedder'").fetchone()
        stored = row[0] if row else None
        if stored is None:
            # first run with this setting: vectors made before tracking came from Gemini
            stored = f"gemini:{config.GEMINI_EMBED_MODEL}:{config.EMBED_DIM}"
            any_vec = con.execute("SELECT 1 FROM chunks WHERE embedding IS NOT NULL LIMIT 1").fetchone()
            if not any_vec:
                stored = current
        if stored != current:
            con.execute("UPDATE chunks SET embedding=NULL")
            con.execute("UPDATE learned SET embedding=NULL")
            cleared = True
        con.execute("INSERT INTO app_settings (key, value) VALUES ('embedder', ?) "
                    "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (current,))
    _space_checked["id"] = current
    return cleared


def semantic_coverage():
    with _conn() as con:
        total = con.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
        with_vec = con.execute("SELECT COUNT(*) FROM chunks WHERE embedding IS NOT NULL").fetchone()[0]
    return with_vec, total


def load_samples_if_empty():
    """Load sample_docs/ when the knowledge base is empty. Region = start of the file name
    (NORAM_, EUROPE_, APAC_, GLOBAL_); type = FAQ / Policy / Quick Guide / SOP by name."""
    with _conn() as con:
        if con.execute("SELECT COUNT(*) FROM documents").fetchone()[0]:
            return
    folder = config.BASE_DIR / "sample_docs"
    if not folder.exists():
        return
    regions = config.REGIONS + [config.GLOBAL_REGION]
    for f in sorted(folder.iterdir()):
        region = f.name.split("_", 1)[0].upper()
        if region not in regions or f.suffix.lstrip(".").lower() not in config.ALLOWED_EXTENSIONS:
            continue
        name = f.name.lower()
        doc_type = ("FAQ" if "faq" in name else "Policy" if "policy" in name
                    else "Quick Guide" if "guide" in name else "SOP")
        title = f.stem.split("_", 1)[-1].replace("_", " ")
        try:
            add_document(region, doc_type, title, f.name, f.read_bytes(), "Sample data", "1.0",
                         sub_lob=_sample_sub_lob(region, name))
        except Exception:
            pass


def _sample_sub_lob(region, name):
    """Sub-LOB tags for the sample documents."""
    subs = get_lobs().get(region, [])
    rules = [("obt", "OBT Cancellation"), ("refund", "Refund"), ("issuance", "Ticketing")]
    for word, sub in rules:
        if word in name and sub in subs:
            return sub
    return ""


# ---------------------------------------------------------------------------
# Pictures inside documents - kept so Vidya can SHOW them with an answer
# (finding answers inside pictures still needs the AI key; showing them does not)
# ---------------------------------------------------------------------------
def _image_location(loc):
    """'Page 3 (scanned)' -> 'Page 3'; 'Refund - Table 2' -> 'Refund'; 'X - image 2' -> ('X', 2)."""
    loc = re.sub(r"\s*\(scanned\)$", "", loc or "").strip()
    m = re.search(r"\s-\simage (\d+)$", loc)
    n = int(m.group(1)) if m else None
    loc = re.sub(r"\s-\s(Table \d+|image \d+)$", "", loc).strip()
    return loc, n


def mask_doc_images(images, with_cards=False):
    """Hide card numbers in a document's pictures (4 at a time). Pictures that can't be checked
    (no Tesseract) are dropped - nothing unchecked is stored.
    with_cards=True also returns [(location, card numbers read)]."""
    from concurrent.futures import ThreadPoolExecutor
    if not getattr(config, "CARD_CHECK", True):
        return (images, []) if with_cards else images
    if not privacy.ocr_ready()[0]:
        return ([], []) if with_cards else []

    def one(item):
        loc, b = item
        if not b or len(b) < config.MIN_IMAGE_BYTES:
            return None
        try:
            png, _, found = privacy.safe_image(b)
            return loc, png, found
        except privacy.PrivacyError:
            return None
    with ThreadPoolExecutor(max_workers=4) as pool:
        done = [x for x in pool.map(one, images) if x]
    out = [(loc, png) for loc, png, _ in done]
    return (out, [(loc, f) for loc, _, f in done if f]) if with_cards else out


def store_images(doc_id, images, masked=True):
    """Save the pictures of one document as PNG files in the upload folder (so backups include them).
    Tiny pictures (icons, logos, lines) are skipped. Returns how many were stored."""
    from PIL import Image
    _delete_images(doc_id)
    if not masked:
        images = mask_doc_images(images)
    stored, counter = 0, {}
    for loc, data in images:
        loc, _ = _image_location(loc)
        counter[loc] = counter.get(loc, 0) + 1            # same numbering as the AI image reader
        if not data or len(data) < config.MIN_IMAGE_BYTES or stored >= config.MAX_STORED_IMAGES_PER_DOC:
            continue
        try:
            im = Image.open(io.BytesIO(data))
            im.load()
            if min(im.size) < config.MIN_IMAGE_SIDE:
                continue
            if im.mode not in ("RGB", "RGBA"):
                im = im.convert("RGBA" if "A" in im.getbands() else "RGB")
            if im.width > config.MAX_IMAGE_WIDTH:
                im = im.resize((config.MAX_IMAGE_WIDTH, int(im.height * config.MAX_IMAGE_WIDTH / im.width)))
            name = f"docimg_{doc_id}_{uuid.uuid4().hex[:8]}.png"
            im.save(config.UPLOAD_DIR / name, "PNG", optimize=True)
        except Exception:                                    # EMF/WMF or damaged pictures cannot be shown
            continue
        with _conn() as con:
            con.execute("INSERT INTO doc_images (doc_id, location, n, filename) VALUES (?,?,?,?)",
                        (doc_id, loc, counter[loc], name))
        stored += 1
    with _conn() as con:
        con.execute("UPDATE documents SET images_done=1 WHERE id=?", (doc_id,))
    return stored


def _delete_images(doc_id):
    with _conn() as con:
        rows = con.execute("SELECT filename FROM doc_images WHERE doc_id=?", (doc_id,)).fetchall()
        con.execute("DELETE FROM doc_images WHERE doc_id=?", (doc_id,))
    for r in rows:
        try:
            (config.UPLOAD_DIR / r["filename"]).unlink(missing_ok=True)
        except OSError:
            pass


def backfill_images():
    """Documents uploaded before pictures were kept: read their stored file once and keep the pictures.
    Safe to call at every start - documents already done are skipped. Returns pictures stored."""
    from pathlib import Path
    with _conn() as con:
        rows = con.execute("SELECT id, filename, stored_path FROM documents "
                           "WHERE COALESCE(images_done,0)=0").fetchall()
    total = 0
    for r in rows:
        try:
            path = Path(r["stored_path"] or "")
            if not path.is_file():
                path = config.UPLOAD_DIR / path.name
            _, images = extract_text(r["filename"], path.read_bytes())
            total += store_images(r["id"], images, masked=False)
        except Exception:
            with _conn() as con:                             # unreadable file: do not retry every start
                con.execute("UPDATE documents SET images_done=1 WHERE id=?", (r["id"],))
    return total


def answer_images(hits, limit=None):
    """Pictures from the document sections the answer came from (best passages first)."""
    limit = config.MAX_ANSWER_IMAGES if limit is None else limit
    out, seen = [], set()
    docs = [h for h in hits if h.get("kind") == "doc" and h.get("doc_id")][:config.IMAGE_SOURCE_PASSAGES]
    with _conn() as con:
        for h in docs:
            loc, n = _image_location(h.get("location", ""))
            rows = con.execute("SELECT id, filename, n FROM doc_images WHERE doc_id=? AND location=? ORDER BY n",
                               (int(h["doc_id"]), loc)).fetchall()
            if n is not None:                                # the AI read this exact picture
                rows = [r for r in rows if r["n"] == n] or rows
            for r in rows:
                path = config.UPLOAD_DIR / r["filename"]
                if r["id"] in seen or not path.is_file():
                    continue
                seen.add(r["id"])
                out.append({"path": str(path), "caption": f"{h['title']} › {loc}"})
                if len(out) >= limit:
                    return out
    return out


def delete_document(doc_id):
    with _conn() as con:
        row = con.execute("SELECT stored_path FROM documents WHERE id=?", (doc_id,)).fetchone()
        con.execute("DELETE FROM chunks WHERE doc_id=?", (doc_id,))
        con.execute("DELETE FROM documents WHERE id=?", (doc_id,))
    _delete_images(doc_id)
    if row:
        try:
            from pathlib import Path
            Path(row["stored_path"]).unlink(missing_ok=True)
        except OSError:
            pass


def set_document_status(doc_id, status):
    """Active documents are searched; Archived ones are kept but ignored."""
    with _conn() as con:
        con.execute("UPDATE documents SET status=? WHERE id=?", (status, doc_id))


def list_documents(region=None, include_global=True, status=None):
    q, args, where = "SELECT * FROM documents", [], []
    if region:
        regions = [region, config.GLOBAL_REGION] if include_global else [region]
        where.append(f"region IN ({','.join('?' * len(regions))})")
        args += regions
    if status:
        where.append("COALESCE(status,'Active')=?")
        args.append(status)
    if where:
        q += " WHERE " + " AND ".join(where)
    q += " ORDER BY region, doc_type, title"
    with _conn() as con:
        return [dict(r) for r in con.execute(q, args)]


def get_document(doc_id):
    with _conn() as con:
        r = con.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
    return dict(r) if r else None


def document_text(doc_id, limit=6000):
    with _conn() as con:
        rows = con.execute("SELECT location, text FROM chunks WHERE doc_id=? ORDER BY id", (doc_id,)).fetchall()
    return "\n\n".join(r["text"] for r in rows)[:limit]


def popular_questions(region, limit=6):
    """Most-asked questions that Vidya answered well - used for 'Try asking'."""
    with _conn() as con:
        rows = con.execute(
            "SELECT question, COUNT(*) n FROM questions WHERE region=? AND result='answered' "
            "AND COALESCE(feedback,0)>=0 AND question NOT LIKE '%[screenshot]' "
            "GROUP BY lower(question) ORDER BY n DESC, MAX(id) DESC LIMIT ?", (region, limit)).fetchall()
    return [r["question"] for r in rows]


def suggested_questions(region, sub_lob=None, limit=6):
    """'Try asking' for the selected LOB tab and sub-LOB chip:
    1) questions staff asked most that Vidya answered well, 2) questions written in the documents
    themselves ('Q: ...' lines, '... - steps' sections). Selected sub-LOB first, then the LOB, then GLOBAL."""
    out = []

    def add(q):
        q = re.sub(r"\s+", " ", q or "").strip()
        if 8 <= len(q) <= 80 and q.lower() not in {x.lower() for x in out}:
            out.append(q)

    with _conn() as con:
        args, extra = [region], ""
        if sub_lob:
            args.append(sub_lob)
            extra = " AND sub_lob=?"
        for r in con.execute(
                "SELECT question, COUNT(*) n FROM questions WHERE region=?" + extra + " AND result='answered' "
                "AND COALESCE(feedback,0)>=0 AND question NOT LIKE '%[screenshot]' "
                "GROUP BY lower(question) ORDER BY n DESC, MAX(id) DESC LIMIT 4", args).fetchall():
            add(r["question"])
        scopes = ([(region, sub_lob)] if sub_lob else []) + [(region, None), (config.GLOBAL_REGION, None)]
        for reg, sl in scopes:
            if len(out) >= limit:
                break
            sql = ("SELECT c.location, c.text FROM chunks c JOIN documents d ON d.id=c.doc_id "
                   "WHERE d.region=? AND COALESCE(d.status,'Active')='Active'")
            a = [reg]
            if sl:
                sql += " AND d.sub_lob=?"
                a.append(sl)
            rows = con.execute(sql + " ORDER BY c.id LIMIT 400", a).fetchall()
            found = []
            for r in rows:
                for q in re.findall(r"(?m)^\s*Q\s*[:.]\s*(.+\?)\s*$", r["text"]):
                    found.append(q)
                loc = r["location"] or ""
                m = re.match(r"(.+?)\s-\s(?:steps|How to (.+?) step by step)$", loc)
                if loc.endswith(" - steps"):
                    found.append(f"What are the steps for {loc[:-8]}?")
                elif re.match(r".+ - How to .+ step by step$", loc):
                    sysname, what = loc.split(" - How to ", 1)
                    found.append(f"How to {what.replace(' step by step', '')} in {sysname}?")
            if len(found) < 3:                               # SOPs without Q&A: use their section titles
                skip = re.compile(r"^(document|scope|purpose|overview|introduction|revision|version|page \d|"
                                  r"sheet |slide |contents|appendix|references|definitions|image|ownership|"
                                  r"approval|change history|how to use)", re.I)
                for r in rows:
                    loc = re.sub(r"\s-\s(Table \d+|image \d+)$", "", r["location"] or "")
                    loc = re.sub(r"^\d+(\.\d+)*[.)]?\s+", "", loc).strip()
                    if 10 <= len(loc) <= 70 and not skip.match(loc) and not loc.endswith(" - steps"):
                        found.append(loc)
            seen = set()
            for q in found:                                  # spread across documents, not 6 from one
                if q not in seen:
                    seen.add(q)
                    add(q)
                if len(out) >= limit:
                    break
    return out[:limit]


def read_file(doc):
    from pathlib import Path
    p = Path(doc["stored_path"])
    return p.read_bytes() if p.exists() else b""


# ===========================================================================
# 4. SEARCH  (Layer 2)
# ===========================================================================

_CACHE = {}          # (regions, library signature) -> items + word/letter indexes (see _lex_index)


def _library_signature():
    """Changes whenever a document, piece, meaning-vector or learned answer changes - so the
    search index is rebuilt only then, not for every question."""
    with _conn() as con:
        return tuple(con.execute(
            "SELECT (SELECT COUNT(*) || ':' || COALESCE(MAX(id), 0) || ':' || COALESCE(SUM(LENGTH(text)), 0) || ':'"
            " || COALESCE(SUM(LENGTH(embedding)), 0) FROM chunks),"
            " (SELECT COUNT(*) || ':' || COALESCE(SUM(id * 3 + LENGTH(COALESCE(status, '')) + LENGTH(COALESCE(title, ''))"
            " + LENGTH(COALESCE(sub_lob, '')) + LENGTH(COALESCE(keywords, '')) + LENGTH(COALESCE(version, ''))"
            " + LENGTH(COALESCE(region, '')) + LENGTH(COALESCE(lob, '')) + LENGTH(COALESCE(effective_date, ''))"
            " + LENGTH(COALESCE(doc_type, ''))), 0) FROM documents),"
            " (SELECT COUNT(*) || ':' || COALESCE(MAX(id), 0) || ':' || COALESCE(SUM(LENGTH(question) + LENGTH(answer)"
            " + LENGTH(COALESCE(variants, '')) + COALESCE(active, 1) * 7 + LENGTH(COALESCE(region, ''))"
            " + LENGTH(COALESCE(source, '')) + LENGTH(COALESCE(embedding, ''))), 0) FROM learned)").fetchone())


def _cached(region, all_regions):
    key = (region, bool(all_regions), str(config.DB_PATH), _library_signature())
    hit = _CACHE.get(key)
    if hit is None:
        items = _load_candidates(region, all_regions)
        hit = {"items": items, "lex": None}
        if len(_CACHE) > 12:
            _CACHE.clear()
        _CACHE[key] = hit
    return hit


def _candidates(region, all_regions=False):
    """Everything searchable for a LOB (copies, so a search can mark them without changing the cache)."""
    return [dict(i) for i in _cached(region, all_regions)["items"]]


def _load_candidates(region, all_regions=False):
    """Everything searchable for a LOB: its active docs + GLOBAL docs + learned answers.
    all_regions=True: every LOB (used only when the LOB itself has no answer)."""
    regions = tuple(config.REGIONS) + (config.GLOBAL_REGION,) if all_regions else (region, config.GLOBAL_REGION)
    marks = ",".join("?" * len(regions))
    with _conn() as con:
        rows = con.execute(
            "SELECT c.id, c.text, c.location, c.embedding, d.id AS doc_id, d.title, d.doc_type, "
            "d.region, d.version, d.keywords, d.lob, d.sub_lob, d.effective_date, d.uploaded_at "
            "FROM chunks c JOIN documents d ON d.id = c.doc_id "
            f"WHERE c.region IN ({marks}) AND COALESCE(d.status,'Active')='Active'", regions).fetchall()
        learned = con.execute(
            "SELECT id, question, answer, added_by, added_at, region, source, embedding, lob, variants FROM learned "
            f"WHERE region IN ({marks}) AND COALESCE(active,1)=1", regions).fetchall()
    items = [
        {"kind": "doc", "text": r["text"], "title": r["title"], "type": r["doc_type"],
         "location": r["location"], "region": r["region"], "version": r["version"],
         "doc": f"d{r['doc_id']}", "doc_id": r["doc_id"], "chunk_id": r["id"],
         "key": f"c{r['id']}", "vec": _blob_to_vec(r["embedding"]),
         # the section heading counts twice: headings are a strong signal
         "search": f"{r['title']} {r['location']} {r['location']} {r['text']}",
         "keywords": r["keywords"] or "", "lob": r["lob"] or "", "sub_lob": r["sub_lob"] or "",
         "effective": r["effective_date"] or "", "uploaded": (r["uploaded_at"] or "")[:10],
         "authority": config.DOC_AUTHORITY.get(r["doc_type"], 1)}
        for r in rows]
    for l in learned:
        auto = (l["source"] or "verified") == "auto"
        items.append(
            {"kind": "learned", "auto": auto, "text": l["answer"],
             "title": "Learned from staff feedback" if auto else "SME / process team answer",
             "type": "Auto-learned Q&A" if auto else "Verified Q&A",
             "location": f"{'learned' if auto else 'added by ' + (l['added_by'] or '')} on {l['added_at']}",
             "region": l["region"], "question": l["question"], "doc": f"l{l['id']}",
             "key": f"l{l['id']}", "vec": _blob_to_vec(l["embedding"]),
             "search": f"{l['question']} {l['question']} {l['variants'] or ''} {l['answer']}",
             "lob": l["lob"] or "", "sub_lob": ""})
    return items


# Small stop-word list on purpose: words like "name", "change", "first" matter in SOPs.
STOP_WORDS = """a an the is are was were be been am do does did i me my we our you your it its
of to in on for at by with from and or as this that these those can could should would will
how what when where which who why please tell know need want get about any there""".split()

_PLURAL = re.compile(r"\b([a-z]{3,}[^s\W])s\b")


def _norm(text):
    """lower-case, '2,000' -> '2000', simple singular form ('commands' -> 'command')."""
    return _PLURAL.sub(r"\1", re.sub(r"(\d),(\d{3})", r"\1\2", text.lower()))


def _feedback_multipliers(region, query):
    """SELF-LEARNING: passages staff marked 👍 / 👎 for SIMILAR questions are
    moved up (x1.3 per vote, max x1.9) or down (x0.5 per vote)."""
    with _conn() as con:
        rows = con.execute("SELECT question, source_key, value FROM feedback_pairs "
                           "WHERE region IN (?,?)", (region, config.GLOBAL_REGION)).fetchall()
    if not rows:
        return {}
    texts = [r["question"] for r in rows]
    vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit(texts + [query])
    sims = cosine_similarity(vec.transform([query]), vec.transform(texts)).ravel()
    mult = {}
    for r, s in zip(rows, sims):
        if s >= config.FEEDBACK_SIMILARITY:
            m = mult.get(r["source_key"], 1.0)
            mult[r["source_key"]] = min(1.9, m * 1.3) if r["value"] > 0 else m * 0.5
    return mult


def lob_factor(item, region, sub_lob=None):
    """Priority: the selected sub-LOB's documents first; other sub-LOBs and general documents of
    the LOB are still searched. Other LOBs (only in the fallback search) rank lower."""
    f = 1.0
    if item.get("region") not in (region, config.GLOBAL_REGION):
        f *= 0.7
    if sub_lob:
        if not item.get("sub_lob"):
            f *= config.GENERAL_DOC_FACTOR       # general / GLOBAL documents: still used, just after the sub-LOB's
        else:
            f *= config.SUB_LOB_BOOST if item["sub_lob"] == sub_lob else config.OTHER_SUB_FACTOR
    return f


_DEF_Q = re.compile(r"^\s*(?:what\s+is|what's|whats|wat\s+is|define|definition\s+of|meaning\s+of|full\s+form\s+of)\s+"
                    r"(?:an?\s+|the\s+)?(.+?)(?:\s+used\s+for|\s+mean|\s+stand\s+for)?\s*\??\s*$", re.I)
_DEF_Q2 = re.compile(r"^\s*what\s+does\s+(.+?)\s+(?:mean|stand\s+for|do)\s*\??\s*$", re.I)


def definition_term(question):
    """'what is AGM?' -> 'agm'; 'define ADM' -> 'adm'. None for other questions (or long terms)."""
    m = _DEF_Q.match(question or "") or _DEF_Q2.match(question or "")
    if not m:
        return None
    term = m.group(1).strip(" ?.").lower()
    # 'what is AGM process / system / tool' -> AGM
    term = re.sub(r"\s+(process|processes|system|tool|application|app|platform|team)$", "", term).strip()
    return term if 0 < len(term.split()) <= 3 else None


# ---- Exact look-ups: "BLR airport", "code for Heathrow", "capital of Denmark", "airlines of Sweden" ----
# Lists of one-line entries (airports, countries, codes) are long; word search finds the LIST, this
# finds the LINE. Used when the question names a thing and asks for one fact about it.
_LOOKUP_INTENT = {"airport": "airport", "airports": "airport", "code": "code", "codes": "code",
                  "capital": "capital", "currency": "currency", "currencies": "currency",
                  "airline": "airline", "airlines": "airline", "carrier": "airline", "carriers": "airline",
                  "station": "station", "stations": "station", "country": "country", "city": "city",
                  "decode": "code", "encode": "code", "iata": "code", "card": "card", "rail": "rail",
                  "train": "rail", "trains": "rail", "tc": "tc", "area": "tc", "zone": "tc"}
_LOOKUP_FILL = set("""which what whats what's is are the a an of for in to from does do fly flies flying
where located main major name names full form stand stands tell me please show give list all its
belongs belong under falls fall""".split())
_INTENT_IN_LINE = {"airport": r"airport|intl|\b[A-Z]{3} - ", "code": r"\b[A-Z0-9]{2,3}\b",
                   "capital": r"capital", "currency": r"currency|\([A-Z]{3}\)", "airline": r"airline",
                   "station": r"station|central|hbf|gare|\bst\b", "country": r".", "city": r"city|\b[A-Z]{3} - ",
                   "card": r"card|\b(VI|CA|AX|DC|JC|TP|DS)\b", "rail": r"rail|train|bahn|sncf|trenitalia|\bsj\b",
                   "tc": r"\bTC ?[123]\b|tc area"}


def _lookup_terms(question):
    """('BLR airport') -> (['BLR'], {'airport'}); None when it is not a look-up question."""
    words = re.findall(r"[A-Za-z0-9][A-Za-z0-9'&.-]*", question or "")
    if not words or len(words) > 9:
        return None
    intents = {_LOOKUP_INTENT[w.lower()] for w in words if w.lower() in _LOOKUP_INTENT}
    terms = [w.strip(".'") for w in words if w.lower() not in _LOOKUP_INTENT and w.lower() not in _LOOKUP_FILL
             and w.lower() not in STOP_WORDS]
    code_only = len(words) <= 2 and any(re.fullmatch(r"[A-Z]{3}", w) for w in words)
    if not terms or len(terms) > 3 or not (intents or code_only):
        return None
    return terms, intents


def _term_rx(t):
    if re.fullmatch(r"[A-Za-z]{2,3}", t):                         # a code: BLR / blr -> exactly BLR
        alts = [re.escape(t.upper())]
    else:
        alts = [re.escape(t)]
    syn = brain.SYNONYMS.get(t.lower(), "")
    if syn:                                                       # amex -> 'american express' or AX
        codes = [w for w in syn.split() if re.fullmatch(r"[A-Z]{2,3}", w)]
        rest = " ".join(w for w in syn.split() if w not in codes)
        alts += ([re.escape(rest)] if rest else []) + [re.escape(c) for c in codes]
    return _CaseMix(alts)


class _CaseMix:
    """Codes match only in capitals (BLR, AX); words / names in any case."""
    def __init__(self, alts):
        self.rx = [re.compile(rf"\b{a}\b") if re.fullmatch(r"[A-Z0-9]{2,3}", a.replace("\\", ""))
                   else re.compile(rf"\b{a}s?\b", re.I) for a in alts]

    def search(self, text):
        return any(r.search(text) for r in self.rx)

    def match(self, text):
        return any(r.match(text) for r in self.rx)


def line_lookup(region, question, sub_lob=None, limit=4):
    """Returns {'lines': [...], 'item': source piece, 'strong': bool} or None."""
    lt = _lookup_terms(question)
    if not lt:
        return None
    terms, intents = lt
    rxs = [_term_rx(t) for t in terms]
    found = []
    for it in _candidates(region):
        if it["kind"] != "doc":
            continue
        loc = it.get("location", "")
        for raw in it["text"].splitlines():
            line = re.sub(r"^\s*(A\s*:\s*|[-•*]\s*)", "", raw).strip()
            if not line or len(line) > 400 or line.lower().startswith("q:"):
                continue
            if not all(r.search(line) or r.search(loc) for r in rxs):
                continue
            if not all(r.search(line) for r in rxs[:1]):
                continue
            s = 1.0
            head = line[:40]
            if any(r.match(head) for r in rxs):                     # the entry FOR this thing: "BLR - ..."
                s += 1.0
            for i in intents:
                if re.search(_INTENT_IN_LINE[i], line, 0 if i in ("code", "tc") else re.I) or i in loc.lower():
                    s += 0.6
                else:
                    s -= 0.4
            strong = s >= 2.0 + 0.6 * len(intents) - 0.01          # entry line + every asked fact in it
            s = s * lob_factor(it, region, sub_lob) - len(line) / 3000
            found.append((s, line, it, strong))
    if not found:
        return None
    found.sort(key=lambda x: -x[0])
    best = found[0][0]
    if best < 1.5:
        return None
    lines, seen = [], set()
    for s, line, it, _ in found:
        if s < best - 0.35 or len(lines) >= limit:
            break
        if line.lower() not in seen:
            seen.add(line.lower())
            lines.append(line)
    return {"lines": lines, "item": found[0][2], "strong": found[0][3], "terms": terms, "intents": intents}


def lookup_satisfied(answer, lk):
    """Does an answer already give what the look-up asks for (the thing AND the fact, at the start)?"""
    head = (answer or "")[:300]
    return (all(_term_rx(t).search(head) for t in lk["terms"])
            and all(re.search(_INTENT_IN_LINE[i], head, 0 if i in ("code", "tc") else re.I)
                    for i in lk["intents"]))


def _defines(location, term):
    """True when a section heading defines the term: 'AGM - ...', 'AGM: ...', 'AGM', 'What is AGM?'."""
    loc = re.sub(r"\s-\s(Table \d+|image \d+)$", "", (location or "").lower()).strip()
    return (loc == term or loc.startswith(term + " - ") or loc.startswith(term + ": ")
            or loc.startswith(term + " (") or f"what is {term}?" in loc or f"what is {term} " in loc + " ")


class _Lex:
    """Word / letter-group match, EXACTLY the numbers of TfidfVectorizer(sublinear_tf=True) fitted on
    the library + the questions - but the slow part (reading every piece of every document) is done
    once and kept until the library changes. Only the question is read per search."""

    def __init__(self, texts, **kw):
        from sklearn.feature_extraction.text import CountVectorizer
        self.cv = CountVectorizer(**kw)
        c = self.cv.fit_transform(texts).tocsr().astype(np.float64)
        self.df = np.bincount(c.indices, minlength=c.shape[1]).astype(np.float64)
        c.data = 1.0 + np.log(c.data)
        self.c, self.n = c, c.shape[0]
        self.vocab = self.cv.vocabulary_
        self.analyze = self.cv.build_analyzer()

    def rows(self, texts):
        x = self.cv.transform(texts).tocsr().astype(np.float64)
        x.data = 1.0 + np.log(x.data)
        return x

    def similarities(self, queries, extra=None):
        """For each question: cosine similarity to every piece (and to the rows of 'extra')."""
        from collections import Counter
        toks = [Counter(self.analyze(q)) for q in queries]
        dfq = np.zeros(len(self.df))
        oov_df = Counter()
        for t in toks:
            for w in t:
                j = self.vocab.get(w)
                if j is None:
                    oov_df[w] += 1
                else:
                    dfq[j] += 1
        big_n = self.n + len(queries)
        idf = np.log((1.0 + big_n) / (1.0 + self.df + dfq)) + 1.0

        def weigh(m):
            m = m.multiply(idf).tocsr()
            norms = np.sqrt(np.asarray(m.multiply(m).sum(axis=1)).ravel())
            norms[norms == 0] = 1.0
            return m, norms
        x, xn = weigh(self.c)
        e, en = weigh(extra) if extra is not None else (None, None)
        out = []
        for t in toks:
            qv = np.zeros(len(self.df))
            sq = 0.0
            for w, tf in t.items():
                wt = 1.0 + np.log(tf)
                j = self.vocab.get(w)
                if j is None:            # a word no document has still counts in the question's length
                    sq += (wt * (np.log((1.0 + big_n) / (1.0 + oov_df[w])) + 1.0)) ** 2
                else:
                    qv[j] = wt * idf[j]
                    sq += qv[j] ** 2
            qn = np.sqrt(sq) or 1.0
            sims = (x @ qv) / xn / qn
            out.append((sims, (e @ qv) / en / qn if e is not None else None))
        return out


def _lex_index(hit, items):
    if hit["lex"] is None:
        texts = [i["search"] + " " + i.get("keywords", "") for i in items]
        heads = [f"{i['title']} {i.get('location', '')} {i.get('question', '')} "
                 + (i["text"].split("\n", 1)[0] if i["text"].lstrip().startswith("Q:") else "") for i in items]
        words = _Lex(texts, stop_words=STOP_WORDS, ngram_range=(1, 2), preprocessor=_norm)
        hit["lex"] = {"words": words, "pieces": _Lex(texts, analyzer="char_wb", ngram_range=(3, 5)),
                      "heads": words.rows(heads)}
    return hit["lex"]


_COMPARE = re.compile(r"\b(difference|differences|compare|comparison|versus|vs\.?)\b", re.I)


def _store_question(region, asked_by, question, answer, sources, status, meta, topic, result, conf, pct,
                    sme_reason, image=None, score=0.0):
    with _conn() as con:
        cur = con.execute(
            "INSERT INTO questions (region, asked_by, question, answer, sources, score, status, asked_at,"
            " meta, topic, result, confidence, gap_status, lob, sub_lob, confidence_pct, sme_reason)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (region, asked_by, question + (" [screenshot]" if image else ""), answer, json.dumps(sources), score,
             status, _now(), json.dumps(meta), topic, result, conf, None if result == "answered" else "Open",
             meta.get("lob", region), meta.get("sub_lob", ""), pct, sme_reason))
        qid = cur.lastrowid
        if image:
            path = config.UPLOAD_DIR / f"question_{qid}.png"
            path.write_bytes(image)
            con.execute("UPDATE questions SET image_path=? WHERE id=?", (str(path), qid))
    return qid


def _fare_rule_answer(region, question, shot_text, asked_by, image, sub_lob, meta, t0):
    """The agent pasted a fare rule (or a screenshot of one): answer from THAT text, not the documents."""
    fr = farerule.answer(question + "\n" + shot_text, from_screenshot=bool(image))
    ask = fr["ask"] or "Fare rule - which penalty applies?"
    determined = fr["determined"]
    conf, pct = ("high", 92) if determined else ("medium", 60)
    if image and not llm.vision_on():
        conf, pct = "medium", min(pct, 75)              # amounts read by text recognition: check on screen
    sme_reason = None if determined else "Fare rule not clear for every case"
    meta.update({"mode": "fare rule", "format": "short", "topic": "Fares", "lob": region, "sub_lob": sub_lob or "",
                 "seconds": round(time.time() - t0, 1), "did_you_mean": "", "documents": 0, "important": "",
                 "confidence_pct": pct, "supporting_docs": 0, "other_lob": "", "images": [],
                 "checks": farerule.CHECKS, "fare_rule": fr["scenario"], "sme_review": bool(sme_reason)})
    meta["layers"].append("fare rule reader")
    meta["notes"].append("Answer read from the fare rule you pasted - not from the SOP documents.")
    stored_q = ask + "\n[fare rule pasted]\n" + (question if not shot_text else shot_text)[:1500]
    qid = _store_question(region, asked_by, stored_q, fr["answer"], [], "needs_review" if sme_reason else "answered",
                          meta, "Fares", "answered", conf, pct, sme_reason, image)
    return {"id": qid, "answer": fr["answer"], "sources": [], "result": "answered", "found": True,
            "confidence": conf, "confidence_pct": pct, "meta": meta, "question": ask, "images": []}


def _parts_missing(r, parts):
    """True when the one answer to the whole question leaves a part uncovered (or nothing was found).
    A question that matches an FAQ / SME answer as a whole, or asks for a comparison, stays whole."""
    if _COMPARE.search(" ".join(parts)):
        return False
    b = r.get("best")
    if r["result"] == "answered" and b and (b["kind"] == "learned" or b.get("faqsim", 0) >= 0.8):
        return False
    if r["result"] != "answered":
        return True
    seen = r["answer"] or ""                     # what the agent would SEE, not the whole passage
    return any(brain.coverage(brain.core_text(p), seen) < 0.75 for p in parts)


def _title(part):
    t = part.strip().rstrip("?").strip()
    return (t[:1].upper() + t[1:]) if t else t


def _combine_parts(subs):
    """One answer with a numbered block per part, its own sources, and a clear 'not found' for any
    part the documents do not cover (never a guessed step)."""
    blocks, hits, notes, pcts, missing, parts, imp = [], [], [], [], [], [], []
    rank = {"high": 3, "medium": 2, "low": 1}
    worst, worst_conf = None, "high"
    for i, (p, x) in enumerate(subs, 1):
        if x["result"] == "answered":
            src = ", ".join(dict.fromkeys(f"{s['title']} - {s['location']}" for s in x["sources"][:2]))
            blocks.append(f"**{i}. {_title(p)}**\n\n{x['answer']}\n\n_Source: {src}_")
            hits += [h for h in x["hits"][:2] if all(h.get("key") != y.get("key") for y in hits)]
            if x["pct"] is not None:
                pcts.append(x["pct"])
            if rank.get(x["conf"], 1) <= rank.get(worst_conf, 3):
                worst, worst_conf = x, x["conf"]
            if x.get("important"):
                imp.append(x["important"])
            parts.append({"part": p, "result": "answered", "pct": x["pct"]})
        else:
            blocks.append(f"**{i}. {_title(p)}**\n\n⚠️ Not found in the documents - please check with "
                          "your supervisor. This part has been sent to the SME Desk.")
            missing.append(p)
            parts.append({"part": p, "result": x["result"], "pct": None})
    notes.append(f"Your question has {len(subs)} parts, so Vidya searched and answered each part separately"
                 + (f" ({len(missing)} not found in the documents)." if missing else "."))
    best = (worst or {}).get("best") or (hits[0] if hits else None)
    return {"answer": "\n\n".join(blocks), "hits": hits, "best": best, "cover": (worst or {}).get("cover", 1.0),
            "conf": "medium" if missing else worst_conf, "mode": (worst or {}).get("mode", "documents"),
            "important": " ".join(imp)[:400], "notes": notes, "pcts": pcts, "missing": missing, "parts": parts}


_GENERIC = set("""today tomorrow yesterday after before first second third right wrong again still other
always never every where which there their about should could would might people things thing""".split())


def faq_question_cover(region, question_core, hit):
    """For an FAQ piece ("Q: ... A: ..."): None when it is a fair match, else the share of the agent's key
    words in the piece, rare words counting more. An FAQ is held back only when the MOST SPECIFIC word of
    the question ("merchant", "reprint", "Germany") is nowhere in it - a common word ("card", "error",
    "blocked") must not make an unrelated FAQ look like the answer. Exact FAQ matches are never checked."""
    text = (hit.get("text") or "").lstrip()
    if hit.get("kind") != "doc" or not text.startswith("Q:") or hit.get("faqsim", 0) >= 0.8:
        return None
    words = list(dict.fromkeys(brain.content_words(question_core)))
    if len(words) < 2:
        return None
    lex = (_CACHE.get(next((k for k in _CACHE if k[0] == region and not k[1]), None)) or {}).get("lex")
    if not lex:
        return None
    ix = lex["words"]

    def weight(w):
        j = ix.vocab.get(_norm(w))
        df = ix.df[j] if j is not None else 0
        return float(np.log((1.0 + ix.n) / (1.0 + df)) + 1.0)
    weights = {w: weight(w) for w in words}
    present = {w for w in words if brain.coverage(w, hit.get("search", text)) >= 1.0}
    specific = [w for w in words if len(w) >= 5 and w not in _GENERIC]   # "next", "day", "today" are not terms
    if not specific:
        return None
    rarest = max(specific, key=lambda w: weights[w])
    if rarest in present:
        return None
    if ix.vocab.get(_norm(rarest)) is None:
        return 0.75          # no document mentions it at all ("Emirates"): a general method may still apply
    return sum(weights[w] for w in present) / sum(weights.values())   # another document is about it


def _doc_label(h):
    bits = [h.get("type") or "document"]
    if h.get("effective"):
        bits.append(f"effective {h['effective']}")
    return ", ".join(bits)


def winner(pair):
    """Two documents disagree: the one with the NEWER effective date wins; with no dates (or the same),
    the higher authority (Policy > SOP > Process Knowledge > FAQ / Quick Guide > Other). None = a real tie,
    which stays a conflict for the SME."""
    a, b = pair
    da, db = a.get("effective") or "", b.get("effective") or ""
    if da and db and da != db:
        return a if da > db else b
    if a.get("authority", 1) != b.get("authority", 1):
        return a if a.get("authority", 1) > b.get("authority", 1) else b
    return None


def search(region, queries, k=config.CANDIDATES, fmt="short", query_vec=None, sub_lob=None, all_regions=False,
           per_doc=None, raw=None):
    """Search ALL documents with every phrasing of the question.
    score = word match (and, with Gemini, meaning match), adjusted by staff feedback
    and by the selected line of business (its documents rank first).
    Keeps at most PER_DOC passages from any one document so answers can be compared."""
    if isinstance(queries, str):
        queries = [queries]
    queries = [q for q in dict.fromkeys(queries) if q and q.strip()]
    hit = _cached(region, all_regions)
    items = [dict(i) for i in hit["items"]]
    if not items or not queries:
        return []
    ix = _lex_index(hit, items)
    wsim = ix["words"].similarities(queries, ix["heads"])
    psim = ix["pieces"].similarities(queries)
    lex = None
    for (w, h), (p, _) in zip(wsim, psim):
        s = 0.6 * w + 0.4 * p + 0.3 * h
        lex = s if lex is None else np.maximum(lex, s)

    sem = None
    if query_vec is not None and any(i["vec"] is not None for i in items):
        sem = np.array([float(np.dot(query_vec, i["vec"])) if i["vec"] is not None
                        and len(i["vec"]) == len(query_vec) else 0.0 for i in items])

    mult = _feedback_multipliers(region, queries[0])
    # SELF-LEARNING: a learned answer whose question is worded like this one wins clearly
    learned_idx = [i for i, it in enumerate(items) if it["kind"] == "learned"]
    if learned_idx:
        qv = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit(
            [items[i]["question"] for i in learned_idx] + queries)
        lq = qv.transform([items[i]["question"] for i in learned_idx])
        qsim = cosine_similarity(qv.transform(queries), lq).max(axis=0)
        for i, sim in zip(learned_idx, qsim):
            items[i]["qsim"] = float(sim)
    # FAQ pieces in documents ("Q: ...") whose question is worded like this one also rank first
    faq_idx = [i for i, it in enumerate(items) if it["kind"] == "doc" and it["text"].lstrip().startswith("Q:")]
    if faq_idx:
        def _qn(t):                              # lower case, no punctuation, no a / an / the
            return re.sub(r"\b(a|an|the)\b", " ", re.sub(r"[^\w\s]", " ", (t or "").lower())).split()
        fq = [" ".join(_qn(re.sub(r"^\s*Q\s*:\s*", "", items[i]["text"].lstrip().split("\n", 1)[0]))) or "-"
              for i in faq_idx]
        # the agent's own wording counts too (search words have synonyms added: 'uk' -> 'united kingdom')
        qq = [" ".join(_qn(q)) or "-" for q in queries + ([raw] if raw else [])]
        fv = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit(fq + qq)
        fsim = cosine_similarity(fv.transform(qq), fv.transform(fq)).max(axis=0)
        for i, sim in zip(faq_idx, fsim):
            items[i]["faqsim"] = float(sim)
    term = definition_term(queries[0])
    for idx, item in enumerate(items):
        boost = 1.0
        if term and _defines(item.get("location", ""), term):
            boost *= config.DEFINITION_BOOST     # "what is AGM" -> the passage that defines AGM, not one that only uses it
        fs = item.get("faqsim", 0)
        if (fs >= config.FAQ_MATCH_SIM and (not sub_lob or item.get("sub_lob") == sub_lob)
                and brain.coverage(raw or queries[0], item["text"]) >= 0.99):   # ...and it covers every key word
            # the agent asked (almost) exactly an FAQ question of the document - the closer, the higher
            boost *= config.FAQ_MATCH_BOOST * (1 + (fs - config.FAQ_MATCH_SIM))
        if item["kind"] == "learned":
            boost = config.AUTO_LEARNED_BOOST if item.get("auto") else config.LEARNED_BOOST
            if item.get("qsim", 0) >= 0.6:
                boost *= 1.6
        if fmt == "table" and "Table" in item.get("location", ""):
            boost *= 1.3                     # asked for a list/table -> prefer tables
        if fmt == "steps" and len(re.findall(r"(?:^|\n)\s*\d+[.)]\s", item["text"])) >= 3:
            boost *= 1.5                     # asked for steps -> prefer real numbered procedures
        elif fmt == "steps" and re.search(r"(^|\s)[2-9][.)]\s", item["text"]):
            boost *= 1.2
        boost *= mult.get(item["key"], 1.0)
        if item["kind"] == "doc":                 # Policy / SOP a little ahead of FAQ / guides
            boost *= 1.0 + 0.015 * (item.get("authority", 1) - 1)
        item["lob_factor"] = lob_factor(item, region, sub_lob)
        boost *= item["lob_factor"]
        item["lex"] = float(lex[idx])
        item["sem"] = float(sem[idx]) if sem is not None else None
        base = item["lex"] if sem is None else 0.7 * item["sem"] + 0.3 * min(1.0, item["lex"] * 2.5)
        item["score"] = base * boost
        item["feedback"] = mult.get(item["key"], 1.0)
    ranked = sorted(items, key=lambda x: x["score"], reverse=True)
    top = ranked[0]["score"] if ranked else 0
    max_per_doc = per_doc or config.PER_DOC
    picked, per_doc = [], {}
    for r in ranked:
        if r["score"] <= 0 or r["score"] < top * 0.35 or len(picked) >= k:
            break
        if per_doc.get(r["doc"], 0) >= max_per_doc:
            continue
        per_doc[r["doc"]] = per_doc.get(r["doc"], 0) + 1
        picked.append(r)
    return picked


def expand_sections(hits, limit=3000):
    """A long section can be split into several pieces. Re-join the pieces that
    directly continue the matched one, so procedures come out complete (all steps).
    FAQ entries (Q: ...) are never merged with their neighbours."""
    out, seen = [], set()
    with _conn() as con:
        for h in hits:
            if h["kind"] != "doc":
                out.append(h); continue
            if h["chunk_id"] in seen:
                continue
            seen.add(h["chunk_id"])
            if h["text"].lstrip().startswith("Q:"):
                out.append(h); continue
            rows = con.execute("SELECT id, text FROM chunks WHERE doc_id=? AND location=? ORDER BY id",
                               (h["doc_id"], h["location"])).fetchall()
            ids = [r["id"] for r in rows]
            if len(rows) < 2 or h["chunk_id"] not in ids:
                out.append(h); continue
            i = j = ids.index(h["chunk_id"])
            while i > 0 and ids[i - 1] == ids[i] - 1 and not rows[i]["text"].lstrip().startswith("Q:"):
                i -= 1
            while j < len(ids) - 1 and ids[j + 1] == ids[j] + 1 and not rows[j + 1]["text"].lstrip().startswith("Q:"):
                j += 1
            lines = []
            for r in rows[i:j + 1]:
                seen.add(r["id"])
                for line in r["text"].splitlines():
                    line = line.strip()
                    if line and not any(line in l or l in line for l in lines[-6:] if len(l) > 20):
                        lines.append(line)
            merged = "\n".join(lines)[:limit]
            # the joined piece is searched/checked with ALL its words, not only the first piece's
            out.append(dict(h, text=merged, search=f"{h['title']} {h['location']} {h['location']} {merged}"))
    return out


def vocabulary(region):
    return brain.build_vocab(i["search"] for i in _candidates(region))


# ===========================================================================
# 5. ANSWERING  (the three layers)
# ===========================================================================

def _public(h):
    return {"title": h["title"], "type": h["type"], "location": h["location"],
            "region": h["region"], "score": round(h["score"], 3), "text": h["text"],
            "key": h["key"], "kind": h["kind"], "doc_id": h.get("doc_id"),
            "lob": h.get("lob", ""), "sub_lob": h.get("sub_lob", "")}


def _variant(hit, names):
    """Which variant (e.g. Amadeus / Sabre) a passage belongs to - by its section name first."""
    loc = hit.get("location", "").lower()
    for n in names:
        if re.search(rf"\b{re.escape(n.lower())}\b", loc):
            return n
    head = hit.get("text", "")[:300].lower()
    found = [n for n in names if re.search(rf"\b{re.escape(n.lower())}\b", head)]
    return found[0] if len(found) == 1 else None


def clarify_options(region, question, queries, sub_lob=None, query_vec=None, anchor=None):
    """When the best answers split between variants the agent did not mention (Amadeus vs Sabre, or two
    sub-LOBs when no sub-LOB is selected), return a short question with buttons instead of guessing."""
    if not config.ASK_CLARIFY or not question.strip():
        return None
    hits = search(region, queries, k=12, query_vec=query_vec, sub_lob=sub_lob, per_doc=12)
    if not hits:
        return None
    st_ = brain.core_text(question)
    top = hits[0]
    if top["lex"] < config.CONF_MEDIUM and (top["sem"] or 0) < config.SEM_MEDIUM:
        return None                                        # weak match: nothing to choose between
    close = [h for h in hits if h["score"] >= top["score"] * config.CLARIFY_CLOSE
             and h["kind"] == "doc" and brain.coverage(st_, h.get("search", h["text"])) >= config.CLARIFY_MIN_COVER]
    if anchor is not None:                                 # the passage Vidya was about to answer from
        close = [anchor] + [h for h in close if h.get("key") != anchor.get("key")]
    ql = " " + question.lower() + " "
    for group, names in config.CLARIFY_GROUPS.items():
        aliases = [a for n in names for a in [n] + config.CLARIFY_ALIASES.get(n, [])]
        if any(re.search(rf"\b{re.escape(a.lower())}\b", ql) for a in aliases):
            continue                                       # the agent already said which one
        if not close or not _variant(close[0], names):
            continue                                       # the best answer is not specific to one variant
        found = []
        for h in close:
            v = _variant(h, names)
            if v and v not in found:
                found.append(v)
        if len(found) >= 2:
            return {"kind": "variant", "group": group,
                    "question": f"Which {group} is this for? The steps are different in each.",
                    "options": [{"label": v, "append": f" in {v}"} for v in found[:4]]}
    if not sub_lob:
        subs = []
        for h in close:
            sl = h.get("sub_lob") or ""
            if sl and h.get("region") != config.GLOBAL_REGION and sl not in subs:
                subs.append(sl)
        if len(subs) >= 2 and not any(s.lower() in ql for s in subs):
            return {"kind": "sub_lob", "group": "sub-LOB",
                    "question": f"Which {region} process is this about? Pick one so I search the right SOP first.",
                    "options": [{"label": s, "sub_lob": s} for s in subs[:4]]}
    return None


def ask(region, question, asked_by, image=None, sub_lob=None, clarify=True, _part=False):
    """question: what the staff member typed (can be empty if they only pasted a screenshot)
    image: bytes of a pasted screenshot, or None. region = the LOB tab, sub_lob = the chip.
    Returns a dict for the screen."""
    t0 = time.time()
    question = privacy.mask((question or "").strip())     # card numbers never saved / sent to AI
    meta = {"layers": [], "notes": []}
    use_ai = llm.ai_on()

    # --- Screenshot: read it first so its text can be searched ---------------
    shot_text, shot_issue = "", ""
    if image:
        try:
            image = privacy.safe_image(image)[0]
        except privacy.PrivacyError as e:
            meta["notes"].append(str(e))
            image = None
    if image:
        if llm.vision_on():
            try:
                shot = brain.read_screenshot(image, question)
                shot_text, shot_issue = shot["text"], shot["issue"]
                meta["layers"].append("read screenshot")
            except llm.AIError as e:
                meta["notes"].append(str(e))
        else:
            ocr = privacy.ocr_text(image)                  # no AI: a fare-rule screenshot is still read (Tesseract)
            if ocr and farerule.detect(ocr):
                shot_text = ocr
                meta["layers"].append("read screenshot (text recognition)")
            else:
                meta["notes"].append("Screenshots are read only by an AI that reads pictures "
                                     "(Gemini/Claude key, or LOCAL_VISION_MODEL for local AI).")

    # --- A pasted fare rule (category 16): read it and say which fee applies -------------------
    if not _part and farerule.detect(question + "\n" + shot_text):
        return _fare_rule_answer(region, question, shot_text, asked_by, image, sub_lob, meta, t0)
    if not question:
        question = shot_issue or "What does this screenshot mean and what should I do?"

    # --- Layer 1: understand ---------------------------------------------------
    fixed, changes = brain.fix_spelling(question, vocabulary(region))
    fmt = brain.detect_format(fixed)
    topic = brain.detect_topic(fixed + " " + shot_text[:200])
    queries = [brain.search_text(question), brain.search_text(fixed)] + ([shot_text[:400]] if shot_text else [])
    meta["layers"].append("spelling + topic")

    qvec = None
    if llm.embeddings_on():
        check_embed_space()
        try:
            qvec = llm.embed([fixed + (" " + shot_issue if shot_issue else "")], query=True)[0]
            meta["layers"].append("semantic search")
        except llm.AIError as e:
            meta["notes"].append(str(e))

    state = {"fixed": fixed, "changes": changes, "fmt": fmt, "queries": queries, "raw": fixed}

    # --- Not clear which GDS / process? Ask first (one click), then answer -------
    # "what is X?" asks for a meaning, which does not depend on the GDS - never ask "Amadeus or Sabre?" for it
    if definition_term(fixed):
        clarify = False
    if clarify and not image:
        opts = clarify_options(region, fixed, queries, sub_lob=sub_lob, query_vec=qvec)
        if opts:
            meta.update({"clarify": opts, "lob": region, "sub_lob": sub_lob or "",
                         "seconds": round(time.time() - t0, 1), "mode": "documents"})
            return {"id": None, "answer": opts["question"], "sources": [], "result": "clarify",
                    "found": False, "confidence": None, "confidence_pct": None, "meta": meta,
                    "question": question, "images": []}

    def find(all_regions):
        """Search (this LOB, or every LOB) and decide: answered / conflict / not_found."""
        fixed, changes, fmt, queries = state["fixed"], state["changes"], state["fmt"], list(state["queries"])
        hits = search(region, queries, raw=state.get("raw"), fmt=fmt, query_vec=qvec, sub_lob=sub_lob, all_regions=all_regions)
        best = hits[0] if hits else None
        weak = (not best) or (best["lex"] < config.CONF_MEDIUM and (best["sem"] or 0) < config.SEM_MEDIUM)
        if use_ai and (weak or image) and qvec is None and not all_regions:
            try:                                # only when the question is unclear - saves quota
                u = brain.ai_understand(question, shot_text)
                queries += [brain.search_text(u["corrected"])] + u["queries"]
                fmt = u["format"] if fmt == "short" else fmt
                if u["corrected"].lower() != question.lower():
                    fixed, changes = u["corrected"], changes or [("", "")]
                hits = search(region, queries, raw=state.get("raw"), fmt=fmt, query_vec=qvec, sub_lob=sub_lob)
                meta["layers"].append("AI rephrase")
                state.update(fixed=fixed, changes=changes, fmt=fmt, queries=queries)
            except llm.AIError as e:
                meta["notes"].append(str(e))

        # --- Layers 2 + 3: compare and answer -----------------------------------
        hits = expand_sections(hits)
        st0 = brain.core_text(fixed)
        hits, used = rerank.apply(fixed, hits, cover=lambda h: brain.coverage(st0, h.get("search", h["text"])))
        if used and "reranker" not in meta["layers"]:
            meta["layers"].append("reranker")
        st = brain.core_text(fixed)
        # re-rank the leaders: a passage that covers more of the question's key words wins
        for h in hits[:5]:
            h["cover"] = brain.coverage(st, h.get("search", h["text"]))
            h["score"] = h["score"] * (0.6 + 0.4 * h["cover"])
        hits = sorted(hits[:5], key=lambda x: x["score"], reverse=True) + hits[5:]
        best = hits[0] if hits else None
        cover = brain.coverage(st, best.get("search", best["text"])) if best else 0
        conf = brain.offline_confidence(best["lex"], cover, best["sem"], len(brain.content_words(st))) if best else "low"
        if best and best["kind"] == "learned":   # short, curated answers: trust a strong match
            if conf == "medium" or (conf == "low" and best["lex"] >= 0.25 and cover >= 0.5):
                conf = "high" if conf == "medium" else "medium"
        out = {"result": "not_found", "answer": None, "important": "", "conflict_pair": None,
               "conflict_text": "", "mode": "documents", "resolved": None}

        if hits and use_ai and (conf != "low" or (best["lex"] >= config.CONF_MEDIUM * 0.6)):
            try:
                a = brain.ai_answer(fixed if changes else question, region, hits, fmt, shot_text,
                                    " › ".join(x for x in (region, sub_lob) if x))
                out["mode"] = "ai"
                meta["layers"].append("AI compare + answer")
                chain, used_ai = llm.providers(), llm.last_used.get("provider")
                if used_ai and chain and used_ai != chain[0]:
                    meta["notes"].append(f"Answered by the backup AI ({llm.name_of(used_ai)}) - "
                                         f"{llm.name_of(chain[0]).split(' (')[0]} is paused for a few minutes.")
                out["result"], conf = a["status"], a["confidence"]
                out["answer"], out["important"], out["conflict_text"] = a["answer"], a["important"], a["conflict"]
                if out["result"] == "answered" and out["answer"]:
                    # every command / number / code must be in the cited document, else the step is marked
                    checked_answer, bad, checked, order = brain.verify_answer(out["answer"], hits)
                    out["answer"] = checked_answer
                    if order:
                        a["used"] = order                  # sources listed in the order the steps cite them
                    if checked:
                        meta["layers"].append("answer checked against documents")
                    if bad:
                        meta["notes"].append(f"{len(bad)} step(s) mention a command or number Vidya could not find "
                                             "in the documents - marked ⚠️. Check them before you act.")
                        if len(bad) * 2 >= checked:        # half the facts unsupported: not a reliable answer
                            conf = {"high": "medium", "medium": "low"}.get(conf, conf)
                        out["unsupported"] = len(bad)
                if a.get("missing") and out["result"] == "answered":
                    out["missing"] = a["missing"]
                    meta["notes"].append(f"Not covered by the documents: {a['missing']}")
                if out["result"] == "answered" and conf == "low":
                    out["result"] = "not_found"
                if a["used"]:                   # show the sources the AI actually used first
                    used = [hits[i - 1] for i in a["used"] if 0 < i <= len(hits)]
                    hits = used + [h for h in hits if h not in used]
                if out["result"] == "conflict":
                    out["conflict_pair"] = [h for h in hits if h["kind"] == "doc"][:2]
            except llm.AIError as e:
                meta["notes"].append(str(e))
                out["mode"] = "documents"

        def documents_decision():
            """The tested word-match rule (no AI). If meaning-search ranked the passages, rank them
            again by words first, so a quota problem can never hide a strong documented answer."""
            h, b, c = hits, best, cover
            if qvec is not None:
                h = expand_sections(search(region, queries, raw=state.get("raw"), fmt=fmt, query_vec=None, sub_lob=sub_lob,
                                           all_regions=all_regions))
                h = rerank.apply(fixed, h, cover=lambda x: brain.coverage(st, x.get("search", x["text"])))[0]
                for x in h[:5]:
                    x["cover"] = brain.coverage(st, x.get("search", x["text"]))
                    x["score"] = x["score"] * (0.6 + 0.4 * x["cover"])
                h = sorted(h[:5], key=lambda x: x["score"], reverse=True) + h[5:]
                b = h[0] if h else None
                c = brain.coverage(st, b.get("search", b["text"])) if b else 0
            if not b:
                return None
            cf = brain.offline_confidence(b["lex"], c, None, len(brain.content_words(st)))
            if b["kind"] == "learned" and (cf == "medium" or (b["lex"] >= 0.25 and c >= 0.5)):
                cf = "high" if cf == "medium" else "medium"
            wq = faq_question_cover(region, st, b)
            if wq is not None:                  # the most specific word of the question is not in this FAQ
                cf = "low" if wq < 0.6 else ("medium" if cf == "high" else cf)
            res = {"hits": h, "best": b, "cover": c, "conf": cf, "result": "not_found",
                   "answer": None, "conflict_pair": None}
            pair = brain.offline_conflict(st, h) if cf != "low" else None
            win = winner(pair) if pair and b["kind"] != "learned" else None
            if win:                            # newer / stronger document decides, the other is flagged
                lose = pair[1] if win is pair[0] else pair[0]
                h = [win] + [x for x in h if x is not win]
                res.update(result="answered", answer=brain.crisp_answer(fixed, h, fmt), hits=h, best=win,
                           conf="medium", resolved=(win, lose))
            elif pair and b["kind"] != "learned":
                res.update(result="conflict", conflict_pair=list(pair))
            elif cf != "low":
                res.update(result="answered", answer=brain.crisp_answer(fixed, h, fmt))
            return res

        if out["mode"] == "documents" and hits:
            d = documents_decision()
            if d:
                hits, best, cover, conf = d["hits"], d["best"], d["cover"], d["conf"]
                out["result"], out["conflict_pair"] = d["result"], d["conflict_pair"]
                out["resolved"] = d.get("resolved")
                if d["answer"]:
                    out["answer"] = d["answer"]
        elif (out["mode"] == "ai" and out["result"] == "not_found" and hits
              and config.DOCUMENT_ANSWER_WHEN_AI_UNSURE):
            # The AI said "not found" (often a cut-short reply on a busy free tier) - if the documents
            # clearly answer the question by the strict word-match rule, show that answer instead.
            d = documents_decision()
            if d and d["result"] == "answered" and d["conf"] in config.DOCUMENT_OVERRIDE_LEVELS:
                hits, best, cover, conf = d["hits"], d["best"], d["cover"], d["conf"]
                out.update(result="answered", answer=d["answer"], mode="documents", important="")
                meta["notes"].append("The AI could not confirm an answer, so this answer is taken directly "
                                     "from the documents - check the source below.")
                meta["layers"].append("document answer (AI unsure)")
        out.update(hits=hits, best=best, cover=cover, conf=conf)
        return out

    r = find(all_regions=False)
    other_lob = ""
    if r["result"] == "not_found" and config.CROSS_LOB_FALLBACK and not image:
        alt = find(all_regions=True)
        top = alt["best"]
        if alt["result"] == "answered" and top and top.get("region") not in (region, config.GLOBAL_REGION):
            other_lob = top["region"]
            alt["conf"] = "medium" if alt["conf"] == "high" else alt["conf"]   # other LOB: never "high"
            meta["notes"].append(f"Not found in {region} documents. This answer comes from {other_lob} "
                                 f"documents - check it applies to {region} before using it.")
            meta["layers"].append("searched other LOBs")
            r = alt
    # --- Exact look-up ("BLR airport", "capital of Denmark"): the one matching LINE of a long list ---
    lk = None if image else line_lookup(region, state["raw"], sub_lob)
    if lk:
        b0 = r["best"]
        curated = bool(b0 and r["result"] == "answered" and (b0["kind"] == "learned" or b0.get("faqsim", 0) >= 0.8))
        ans = (r["answer"] or "").lower()
        use = (r["result"] != "answered") or (
            r["mode"] == "documents" and not curated and lk["lines"][0].lower()[:80] not in ans
            and (lk["strong"] or not lookup_satisfied(r["answer"], lk)))
        if use and (r["mode"] == "documents" or r["result"] != "answered"):
            text = "\n".join(lk["lines"])
            top = dict(lk["item"], text=text, lex=max(lk["item"].get("lex") or 0, config.CONF_HIGH),
                       score=max(lk["item"].get("score") or 0, (b0 or {}).get("score", 0) * 1.01), cover=1.0)
            r.update(result="answered", answer=text, mode="documents", important="", conflict_pair=None,
                     hits=[top] + [h for h in r["hits"] if h.get("key") != top.get("key")],
                     best=top, cover=1.0, conf="high" if lk["strong"] else "medium")
            meta["layers"].append("exact line look-up")
    # --- Several questions in one ("unused ticket, fare changed and invoice wrong"): answer each part ---
    multi = None
    if not _part and not image and getattr(config, "MULTI_PART", True):
        parts = brain.split_parts(question)          # the agent's own words (each part is spell-fixed itself)
        if len(parts) >= 2 and _parts_missing(r, parts):
            subs = [(p, ask(region, p, asked_by, sub_lob=sub_lob, clarify=False, _part=True)) for p in parts]
            if sum(1 for _, x in subs if x["result"] == "answered") >= 1:
                multi = _combine_parts(subs)
                r.update(result="answered", answer=multi["answer"], hits=multi["hits"], best=multi["best"],
                         cover=multi["cover"], conf=multi["conf"], mode=multi["mode"], important=multi["important"],
                         conflict_pair=None, conflict_text="", resolved=None)
                meta["layers"].append(f"answered in {len(parts)} parts")
                meta["parts"] = multi["parts"]
                meta["notes"] += multi["notes"]
    result, answer, important = r["result"], r["answer"], r["important"]
    conflict_pair, conflict_text, mode = r["conflict_pair"], r["conflict_text"], r["mode"]
    hits, best, cover, conf = r["hits"], r["best"], r["cover"], r["conf"]
    fixed, changes, fmt = state["fixed"], state["changes"], state["fmt"]

    # the answer came from e.g. an Amadeus section although the agent did not say Amadeus: ask first
    if clarify and not image and not multi and result == "answered" and best and best["kind"] == "doc":
        opts = clarify_options(region, fixed, state["queries"], sub_lob=sub_lob, query_vec=qvec, anchor=best)
        if opts and opts["kind"] == "variant":
            meta.update({"clarify": opts, "lob": region, "sub_lob": sub_lob or "",
                         "seconds": round(time.time() - t0, 1), "mode": "documents"})
            return {"id": None, "answer": opts["question"], "sources": [], "result": "clarify",
                    "found": False, "confidence": None, "confidence_pct": None, "meta": meta,
                    "question": question, "images": []}

    if result == "not_found":
        answer = config.FALLBACK_TEXT
    elif result == "conflict":
        answer = config.CONFLICT_TEXT

    sources = [_public(h) for h in (conflict_pair or hits)] if result != "not_found" else []
    images = answer_images(hits) if result == "answered" and config.SHOW_ANSWER_IMAGES else []
    support = brain.supporting_documents(hits) if result == "answered" else 0
    pct = brain.confidence_percent(conf, best, cover, support) if best and result == "answered" else None
    if multi and multi["pcts"]:
        pct = min(multi["pcts"])                       # a multi-part answer is as sure as its weakest part
    if _part:                                           # one part of a bigger question: nothing is saved
        return {"result": result, "answer": answer, "sources": sources, "hits": hits, "best": best,
                "cover": cover, "conf": conf, "pct": pct, "mode": mode, "important": important,
                "notes": meta["notes"]}
    meta.update({"mode": mode, "format": fmt, "topic": topic, "seconds": round(time.time() - t0, 1),
                 "did_you_mean": fixed if (changes and fixed.lower() != question.lower()) else "",
                 "documents": len({s["title"] for s in sources}), "important": important,
                 "conflict_text": conflict_text, "screenshot_text": shot_text[:600],
                 "learned": bool(best and best["kind"] == "learned" and result == "answered"),
                 "confidence_pct": pct, "supporting_docs": support, "lob": region, "sub_lob": sub_lob or "",
                 "other_lob": other_lob, "images": images,
                 "sub_lob_match": bool(best and sub_lob and best.get("sub_lob") == sub_lob)})
    # Which questions go to the SME Desk: no answer, documents disagree, or confidence below 70 %.
    # An answer that already comes from an SME / process-team answer is not sent again.
    from_verified = bool(best and best["kind"] == "learned" and not best.get("auto") and result == "answered")
    sme_reason = None
    resolved = r.get("resolved") if result == "answered" else None
    if resolved:
        win, lose = resolved
        meta["notes"].append(f"Two documents disagree. This answer follows **{win['title']}** ({_doc_label(win)}), "
                             f"which is newer or has higher priority than **{lose['title']}** ({_doc_label(lose)}). "
                             "The question went to the SME Desk so the older document can be fixed.")
        meta["resolved"] = {"used": win["title"], "other": lose["title"]}
    if result == "not_found":
        sme_reason = "No answer"
    elif multi and multi["missing"]:
        sme_reason = "Partly answered - nothing found for: " + "; ".join(multi["missing"])[:200]
    elif r.get("missing") and result == "answered":
        sme_reason = ("Partly answered - not in the documents: " + r["missing"])[:240]
    elif r.get("unsupported") and result == "answered":
        sme_reason = "AI answer has steps not found in the documents"
    elif resolved:
        sme_reason = "Documents disagree (newer / higher priority one used)"
    elif result == "conflict":
        sme_reason = "Documents disagree"
    elif pct is not None and pct < config.SME_REVIEW_BELOW_PCT and not from_verified:
        sme_reason = f"Low confidence ({pct}%)"
    meta["sme_review"] = bool(sme_reason) and result == "answered"
    if farerule.is_penalty_question(question):
        meta["checks"] = farerule.CHECKS                # the three checks for every fare-penalty question
    status = "needs_review" if sme_reason else "answered"
    with _conn() as con:
        cur = con.execute(
            "INSERT INTO questions (region, asked_by, question, answer, sources, score, status, asked_at,"
            " meta, topic, result, confidence, gap_status, lob, sub_lob, confidence_pct, sme_reason)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (region, asked_by, question + (" [screenshot]" if image else ""), answer,
             json.dumps(sources), best["score"] if best else 0.0, status, _now(), json.dumps(meta),
             topic, result, conf, None if result == "answered" else "Open", meta["lob"], meta["sub_lob"], pct,
             sme_reason))
        qid = cur.lastrowid
        if image:
            path = config.UPLOAD_DIR / f"question_{qid}.png"
            path.write_bytes(image)
            con.execute("UPDATE questions SET image_path=? WHERE id=?", (str(path), qid))
    return {"id": qid, "answer": answer, "sources": sources, "result": result,
            "found": result == "answered", "confidence": conf, "confidence_pct": pct,
            "meta": meta, "question": question, "images": images}


# ===========================================================================
# 6. LEARNING LOOP
# ===========================================================================

def give_feedback(qid, helpful: bool, comment="", trusted=False):
    """👍 / 👎 from staff. This is where Vidya learns by itself."""
    with _conn() as con:
        q = con.execute("SELECT * FROM questions WHERE id=?", (qid,)).fetchone()
        if not q:
            return
        con.execute("UPDATE questions SET feedback=?, comment=? WHERE id=?",
                    (1 if helpful else -1, privacy.mask(comment), qid))
        if not helpful:
            con.execute("UPDATE questions SET status='needs_review', gap_status=COALESCE(gap_status,'Open'), "
                        "sme_reason='Not helpful 👎' WHERE id=?", (qid,))
        sources = json.loads(q["sources"] or "[]")
        question = re.sub(r"\s*\[screenshot\]$", "", q["question"])
        for s in sources[:2]:              # the passages that produced the answer
            if s.get("key"):
                con.execute("INSERT INTO feedback_pairs (region, question, source_key, value, at) "
                            "VALUES (?,?,?,?,?)", (q["region"], question, s["key"],
                                                   1 if helpful else -1, _now()))
        if not helpful:                    # a wrong auto-learned answer is switched off
            for s in sources[:1]:
                if s.get("key", "").startswith("l") and s.get("type") == "Auto-learned Q&A":
                    con.execute("UPDATE learned SET active=0 WHERE id=? AND source='auto'",
                                (int(s["key"][1:]),))
    if helpful and config.AUTO_LEARN and q["result"] == "answered":
        _auto_learn(q, question, trusted)


def _auto_learn(q, question, trusted=False):
    """Save an answer staff found helpful so similar questions get it straight away."""
    if not sources_from_documents(q):
        return                             # already a learned answer - nothing new to save
    with _conn() as con:
        voters = con.execute("SELECT COUNT(DISTINCT asked_by) FROM questions WHERE region=? AND feedback=1 "
                             "AND lower(question)=lower(?) AND lower(answer)=lower(?)",
                             (q["region"], q["question"], q["answer"])).fetchone()[0]
        if voters < config.AUTO_LEARN_MIN_VOTES and not trusted:
            return
        same = con.execute("SELECT id FROM learned WHERE region=? AND lower(question)=lower(?)",
                           (q["region"], question)).fetchone()
        if same:
            con.execute("UPDATE learned SET votes=COALESCE(votes,0)+1, active=1 WHERE id=?", (same["id"],))
            return
    vec = None
    if llm.embeddings_on():
        try:
            vec = _vec_to_blob(llm.embed([f"{question}\n{q['answer']}"])[0])
        except llm.AIError:
            pass
    with _conn() as con:
        con.execute("INSERT INTO learned (region, question, answer, added_by, added_at, from_question,"
                    " source, active, votes, embedding) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (q["region"], question, q["answer"], "Vidya (staff 👍)", _now(), q["id"],
                     "auto", 1, 1, vec))


def sources_from_documents(q):
    srcs = json.loads(q["sources"] or "[]")
    return bool(srcs) and srcs[0].get("kind") == "doc"


def review_queue(region=None):
    q = "SELECT * FROM questions WHERE status='needs_review'"
    args = []
    if region:
        q += " AND region=?"
        args.append(region)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def teach(region, question, answer, admin, from_question=None, variants=None):
    """Store a VERIFIED answer (written by the process team / an SME). Used first from now on.
    variants = the wordings agents actually typed, so the same words find this answer again."""
    question, answer = privacy.mask(question), privacy.mask(answer)
    variants = [privacy.mask(v) for v in (variants or [])]
    vtext = " | ".join(dict.fromkeys(v.strip() for v in (variants or []) if v and v.strip()
                                     and _norm_q(v) != _norm_q(question)))[:1000]
    vec = None
    if llm.embeddings_on():
        try:
            vec = _vec_to_blob(llm.embed([f"{question}\n{vtext}\n{answer}"])[0])
        except llm.AIError:
            pass
    with _conn() as con:
        con.execute("INSERT INTO learned (region, question, answer, added_by, added_at, from_question,"
                    " source, active, votes, embedding, variants) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                    (region, question, answer, admin, _now(), from_question, "verified", 1, 0, vec, vtext))
        if from_question:
            con.execute("UPDATE questions SET status='resolved', gap_status='Fixed' WHERE id=?",
                        (from_question,))


# ===========================================================================
# 6b. SME DESK - questions an SME should answer (no SOP or document is changed)
# ===========================================================================

def _norm_q(text):
    text = re.sub(r"\s*\[screenshot\]$", "", text or "").lower()
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text)).strip()


def _reason_of(r):
    if r.get("sme_reason"):
        return r["sme_reason"]
    if r.get("feedback") == -1:
        return "Not helpful 👎"
    return {"not_found": "No answer", "conflict": "Documents disagree"}.get(r.get("result"), "Needs review")


def sme_queue(region=None):
    """Waiting questions, the same question asked several times shown once. Most-asked first."""
    groups = {}
    for r in review_queue(region):
        key = (r["region"], _norm_q(r["question"]))
        g = groups.get(key)
        if g is None:
            g = groups[key] = {"region": r["region"], "question": re.sub(r"\s*\[screenshot\]$", "", r["question"]),
                               "ids": [], "count": 0, "users": [], "reasons": [], "latest": r,
                               "comments": [], "first": r["asked_at"], "last": r["asked_at"]}
        g["ids"].append(r["id"])
        g["count"] += 1
        g["first"] = min(g["first"], r["asked_at"])
        g["last"] = max(g["last"], r["asked_at"])
        if r["asked_by"] not in g["users"]:
            g["users"].append(r["asked_by"])
        reason = _reason_of(r)
        if reason not in g["reasons"]:
            g["reasons"].append(reason)
        if r.get("comment"):
            g["comments"].append(r["comment"])
    return sorted(groups.values(), key=lambda g: (g["count"], g["last"]), reverse=True)


def sme_answer(ids, region, question, answer, sme):
    """Save the SME's answer as a verified answer (searched first next time) and close the questions.
    Uploaded SOPs / documents are NOT changed."""
    question, answer = (question or "").strip(), (answer or "").strip()
    if not question or not answer:
        raise ValueError("Please write the question and the answer.")
    asked = []
    if ids:
        with _conn() as con:
            asked = [re.sub(r"\s*\[screenshot\]$", "", r[0]) for r in con.execute(
                f"SELECT question FROM questions WHERE id IN ({','.join('?' * len(ids))})", ids)]
    teach(region, question, answer, sme, from_question=ids[0] if ids else None, variants=asked)
    _close(ids, sme, "resolved", "Fixed")


def _close(ids, sme, status, gap):
    if not ids:
        return
    marks = ",".join("?" * len(ids))
    with _conn() as con:
        con.execute(f"UPDATE questions SET status=?, gap_status=CASE WHEN gap_status IS NULL THEN NULL ELSE ? END,"
                    f" sme_by=?, sme_at=? WHERE id IN ({marks})", [status, gap, sme, _now(), *ids])


def sme_confirm(ids, sme):
    """The SME checked Vidya's answer and it is right: keep it as a verified answer."""
    with _conn() as con:
        q = con.execute("SELECT * FROM questions WHERE id=?", (ids[0],)).fetchone()
    if not q or q["result"] != "answered":
        raise ValueError("There is no Vidya answer to confirm - please write the answer.")
    sme_answer(ids, q["region"], re.sub(r"\s*\[screenshot\]$", "", q["question"]), q["answer"], sme)


def sme_flag_sop(ids, sme, note=""):
    """Needs a change in the SOP itself: goes to Admin → Knowledge Gaps."""
    set_gap(ids, status="SOP update needed", owner="", note=note or f"Flagged by SME {sme}")
    _close(ids, sme, "sop_update", "SOP update needed")


def sme_dismiss(ids, sme):
    _close(ids, sme, "dismissed", "Closed")


def update_learned(lid, question, answer, editor):
    question, answer = privacy.mask((question or "").strip()), privacy.mask((answer or "").strip())
    if not question or not answer:
        raise ValueError("Question and answer can't be empty.")
    vec = None
    if llm.embeddings_on():
        try:
            vec = _vec_to_blob(llm.embed([f"{question}\n{answer}"])[0])
        except llm.AIError:
            pass
    with _conn() as con:   # the agents' wordings (variants) are kept
        con.execute("UPDATE learned SET question=?, answer=?, added_by=?, added_at=?, embedding=? WHERE id=?",
                    (question, answer, editor, _now(), vec, lid))


def sme_answers(region=None):
    """Verified answers (written by SMEs / process team), newest first."""
    return list_learned(region, "verified")


# ===========================================================================
# 6c. LIVE CHAT WITH AN SME - every message is kept; the SME chooses which question/answer
#     pairs become verified answers (used first next time). Works on any server: the pages
#     refresh themselves every few seconds and read the same database.
# ===========================================================================
CHAT_MAX_LEN = 2000
ONLINE_SECONDS = 45          # an SME counts as online if the SME Desk was open in the last 45 s


def set_presence(user, region, available=True):
    with _conn() as con:
        con.execute("INSERT INTO sme_presence (user, region, available, last_seen) VALUES (?,?,?,?) "
                    "ON CONFLICT(user) DO UPDATE SET region=excluded.region, available=excluded.available, "
                    "last_seen=excluded.last_seen", (user, region, 1 if available else 0, time.time()))


def smes_online(region=None):
    q, args = "SELECT user, region FROM sme_presence WHERE available=1 AND last_seen>=?", [time.time() - ONLINE_SECONDS]
    if region:
        q += " AND region=?"
        args.append(region)
    with _conn() as con:
        return [dict(r) for r in con.execute(q, args)]


def _msg(con, chat_id, sender, role, text, image_path=None):
    con.execute("INSERT INTO sme_messages (chat_id, sender, role, text, at, image_path) VALUES (?,?,?,?,?,?)",
                (chat_id, sender, role, privacy.mask(text), _now(), image_path))


def _chat_image(chat_id, image):
    """Hide card numbers, then save. Raises privacy.PrivacyError - the picture is then NOT saved."""
    png = privacy.safe_image(image)[0]
    path = config.UPLOAD_DIR / f"chat_{chat_id}_{uuid.uuid4().hex[:8]}.png"
    path.write_bytes(png)
    return str(path)


def start_chat(region, agent, text, sub_lob="", from_question=None, image=None):
    text = (text or "").strip()[:CHAT_MAX_LEN]
    if not text:
        raise ValueError("Please type your question.")
    with _conn() as con:
        cur = con.execute("INSERT INTO sme_chats (region, sub_lob, agent, status, started_at, from_question) "
                          "VALUES (?,?,?,?,?,?)", (region, sub_lob or "", agent, "waiting", _now(), from_question))
        cid = cur.lastrowid
    path = _chat_image(cid, image) if image else None
    with _conn() as con:
        _msg(con, cid, agent, "agent", text, path)
    return cid


def get_chat(chat_id):
    with _conn() as con:
        r = con.execute("SELECT * FROM sme_chats WHERE id=?", (chat_id,)).fetchone()
    return dict(r) if r else None


def chat_messages(chat_id):
    with _conn() as con:
        return [dict(r) for r in con.execute("SELECT * FROM sme_messages WHERE chat_id=? ORDER BY id", (chat_id,))]


def send_chat(chat_id, sender, role, text, image=None):
    """image = screenshot bytes; card numbers are hidden before it is saved (refused if it can't be checked)."""
    text = (text or "").strip()[:CHAT_MAX_LEN]
    c = get_chat(chat_id)
    if (not text and not image) or not c or c["status"] == "closed":
        return False
    path = _chat_image(chat_id, image) if image else None
    with _conn() as con:
        _msg(con, chat_id, sender, role, text or "📷 Screenshot", path)
    return True


def agent_open_chat(agent):
    with _conn() as con:
        r = con.execute("SELECT * FROM sme_chats WHERE agent=? AND status IN ('waiting','active') "
                        "ORDER BY id DESC LIMIT 1", (agent,)).fetchone()
    return dict(r) if r else None


def waiting_chats(region=None):
    q, args = "SELECT * FROM sme_chats WHERE status='waiting'", []
    if region:
        q += " AND region=?"
        args.append(region)
    with _conn() as con:
        rows = [dict(r) for r in con.execute(q + " ORDER BY id", args)]
        for r in rows:
            m = con.execute("SELECT text FROM sme_messages WHERE chat_id=? AND role='agent' ORDER BY id LIMIT 1",
                            (r["id"],)).fetchone()
            r["first"] = m[0] if m else ""
    return rows


def sme_active_chats(sme):
    with _conn() as con:
        return [dict(r) for r in con.execute("SELECT * FROM sme_chats WHERE sme=? AND status='active' ORDER BY id",
                                             (sme,))]


def pick_up_chat(chat_id, sme):
    """Only one SME gets the chat (the first to click)."""
    with _conn() as con:
        cur = con.execute("UPDATE sme_chats SET status='active', sme=?, picked_at=? WHERE id=? AND status='waiting'",
                          (sme, _now(), chat_id))
        if cur.rowcount:
            _msg(con, chat_id, "Vidya", "system", f"{sme} (SME) joined the chat.")
        return bool(cur.rowcount)


def close_chat(chat_id, by):
    c = get_chat(chat_id)
    if not c or c["status"] == "closed":
        return
    with _conn() as con:
        con.execute("UPDATE sme_chats SET status='closed', closed_at=?, closed_by=? WHERE id=?", (_now(), by, chat_id))
        _msg(con, chat_id, "Vidya", "system", f"Chat ended by {by}.")
        replied = con.execute("SELECT COUNT(*) FROM sme_messages WHERE chat_id=? AND role='sme'", (chat_id,)).fetchone()[0]
        if not replied:
            # nobody answered: the question goes to the SME Desk so it is not lost
            con.execute("UPDATE sme_chats SET saved=1 WHERE id=?", (chat_id,))
            if not c.get("from_question"):
                first = con.execute("SELECT text FROM sme_messages WHERE chat_id=? AND role='agent' ORDER BY id LIMIT 1",
                                    (chat_id,)).fetchone()
                if first:
                    con.execute("INSERT INTO questions (region, asked_by, question, answer, sources, score, status,"
                                " asked_at, meta, topic, result, gap_status, lob, sub_lob, sme_reason)"
                                " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                                (c["region"], c["agent"], first[0], "", "[]", 0.0, "needs_review", _now(), "{}",
                                 "", "not_found", "Open", c["region"], c["sub_lob"], "Live chat - no SME reply"))


_CHATTER = re.compile(r"^(hi|hello|hey|ok|okay|thanks|thank you|thx|ty|yes|no|sure|wait|one (min|moment)|"
                      r"pls wait|please wait|welcome|np|noted|done|got it)\b[\s!.]*$", re.I)


def chat_pairs(chat_id):
    """Suggested question/answer pairs: the agent's message(s), then the SME's reply(ies).
    Greetings and 'ok / thanks' are left out. The SME edits them before saving."""
    pairs, q, a = [], [], []
    for m in chat_messages(chat_id):
        if m["role"] == "system" or _CHATTER.match(m["text"].strip()) or m["text"] == "📷 Screenshot":
            continue
        if m["role"] == "agent":
            if a:
                pairs.append((q, a))
                q, a = [], []
            q.append(m["text"])
        elif m["role"] == "sme" and q:
            a.append(m["text"])
    if q and a:
        pairs.append((q, a))
    return [{"question": " ".join(x).strip(), "asked": x, "answer": "\n".join(y).strip()}
            for x, y in pairs if len(" ".join(y)) >= 15]


def save_chat_pairs(chat_id, pairs, sme):
    """Save the chosen pairs as verified answers and close the chat review."""
    c = get_chat(chat_id)
    if not c:
        return 0
    n = 0
    for p in pairs:
        q, a = (p.get("question") or "").strip(), (p.get("answer") or "").strip()
        if q and a:
            teach(c["region"], q, a, sme, from_question=c.get("from_question"), variants=p.get("asked") or [])
            n += 1
    with _conn() as con:
        con.execute("UPDATE sme_chats SET saved=1, saved_count=saved_count+? WHERE id=?", (n, chat_id))
    if n and c.get("from_question"):
        _close([c["from_question"]], sme, "resolved", "Fixed")
    return n


def chats_to_review(sme=None, region=None):
    """Ended chats whose answers nobody has saved yet."""
    q, args = "SELECT * FROM sme_chats WHERE status='closed' AND COALESCE(saved,0)=0", []
    if sme:
        q += " AND sme=?"
        args.append(sme)
    if region:
        q += " AND region=?"
        args.append(region)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def chat_history(region=None, search="", limit=100):
    q, args = "SELECT * FROM sme_chats WHERE 1=1", []
    if region:
        q += " AND region=?"
        args.append(region)
    with _conn() as con:
        chats = [dict(r) for r in con.execute(q + " ORDER BY id DESC LIMIT ?", args + [limit * 5])]
    out = []
    for c in chats:
        msgs = chat_messages(c["id"])
        if search and not any(search.lower() in m["text"].lower() for m in msgs):
            continue
        c["messages"] = msgs
        out.append(c)
        if len(out) >= limit:
            break
    return out


def chats_excel(region=None):
    import pandas as pd
    rows = []
    for c in chat_history(region, limit=5000):
        for m in c["messages"]:
            rows.append({"Chat": c["id"], "LOB": c["region"], "Sub-LOB": c["sub_lob"], "Agent": c["agent"],
                         "SME": c["sme"] or "", "Status": c["status"], "Time": m["at"], "From": m["sender"],
                         "Role": m["role"], "Message": m["text"]})
    buf = io.BytesIO()
    pd.DataFrame(rows or [{"Chat": "", "Message": "no chats yet"}]).to_excel(buf, index=False, sheet_name="SME chats")
    return buf.getvalue()


def mask_existing_data(images=True):
    """One-time clean-up for data saved before card masking existed: masks card numbers in all saved
    text, and in saved screenshots / document pictures (when Tesseract is available).
    Safe to call at every start - it runs once (and the picture part once Tesseract is there)."""
    from pathlib import Path
    with _conn() as con:
        done = {r[0] for r in con.execute("SELECT key FROM app_settings WHERE key LIKE 'card_mask_%'")}
    fixed = 0
    rule = f"_{privacy.min_digits()}"               # runs again when the masking rule changes (v5.1: 13 digits)
    tkey, ikey = "card_mask_text" + rule, "card_mask_images" + rule
    if tkey not in done:
        cols = {"questions": ["question", "answer", "comment"], "chunks": ["text"],
                "learned": ["question", "answer", "variants"], "sme_messages": ["text"],
                "announcements": ["title", "body"]}
        with _conn() as con:
            for table, names in cols.items():
                for r in con.execute(f"SELECT id, {', '.join(names)} FROM {table}").fetchall():
                    new = [privacy.mask(r[n]) if r[n] else r[n] for n in names]
                    if new != [r[n] for n in names]:
                        con.execute(f"UPDATE {table} SET {', '.join(n + '=?' for n in names)} WHERE id=?",
                                    (*new, r["id"]))
                        fixed += 1
            con.execute("INSERT OR REPLACE INTO app_settings (key, value) VALUES (?, ?)", (tkey, _now()))
    if images and ikey not in done and privacy.ocr_ready()[0]:
        files = []
        with _conn() as con:
            files += [r[0] for r in con.execute("SELECT image_path FROM questions WHERE image_path IS NOT NULL")]
            files += [r[0] for r in con.execute("SELECT image_path FROM sme_messages WHERE image_path IS NOT NULL")]
            files += [str(config.UPLOAD_DIR / r[0]) for r in con.execute("SELECT filename FROM doc_images")]
        for f in files:
            p = Path(f or "")
            if p.is_file():
                try:
                    png, n, _ = privacy.safe_image(p.read_bytes())
                    if n:
                        p.write_bytes(png)
                        fixed += 1
                except privacy.PrivacyError:
                    pass
        with _conn() as con:
            con.execute("INSERT OR REPLACE INTO app_settings (key, value) VALUES (?, ?)", (ikey, _now()))
    return fixed


def dismiss(qid):
    with _conn() as con:
        con.execute("UPDATE questions SET status='dismissed', gap_status='Closed' WHERE id=?", (qid,))


def list_learned(region=None, source=None):
    q, args, where = "SELECT * FROM learned", [], []
    if region:
        where.append("region=?"); args.append(region)
    if source:
        where.append("COALESCE(source,'verified')=?"); args.append(source)
    if where:
        q += " WHERE " + " AND ".join(where)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def set_learned(lid, active=None, delete=False, make_verified=False, admin=""):
    with _conn() as con:
        if delete:
            con.execute("DELETE FROM learned WHERE id=?", (lid,))
        if active is not None:
            con.execute("UPDATE learned SET active=? WHERE id=?", (1 if active else 0, lid))
        if make_verified:
            con.execute("UPDATE learned SET source='verified', added_by=? WHERE id=?", (admin, lid))


# ===========================================================================
# 7. KNOWLEDGE GAPS  (questions Vidya could not answer well)
# ===========================================================================

GAP_STATUSES = ["Open", "SOP update needed", "In progress", "Possibly answered", "Fixed", "Closed"]


def _gap_rows(region=None, since=None):
    q = ("SELECT * FROM questions WHERE (result IN ('not_found','conflict') OR feedback=-1 "
         "OR (result IS NULL AND status='needs_review') OR gap_status='SOP update needed')")
    args = []
    if region:
        q += " AND region=?"; args.append(region)
    if since:
        q += " AND asked_at>=?"; args.append(since)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def knowledge_gaps(region=None, since=None, include_closed=False):
    """Group similar unanswered / conflicting / 👎 questions so the process team
    can see what is missing and fix the SOPs. Most-asked first."""
    rows = _gap_rows(region, since)
    if not include_closed:
        rows = [r for r in rows if (r.get("gap_status") or "Open") not in ("Fixed", "Closed")]
    if not rows:
        return []
    texts = [re.sub(r"\s*\[screenshot\]$", "", r["question"]) for r in rows]
    vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit(texts)
    m = vec.transform(texts)
    sims = cosine_similarity(m)
    groups, used = [], set()
    for i in range(len(rows)):
        if i in used:
            continue
        members = [j for j in range(len(rows)) if j not in used and (j == i or sims[i, j] >= 0.5)]
        used.update(members)
        g = [rows[j] for j in members]
        kinds = []
        for r in g:
            kinds.append("Conflict" if r.get("result") == "conflict" else
                         "Not helpful" if r.get("feedback") == -1 else "No answer")
        groups.append({
            "question": texts[i], "count": len(g), "ids": [r["id"] for r in g],
            "regions": sorted({r["region"] for r in g}), "topic": g[0].get("topic") or "Other",
            "type": max(set(kinds), key=kinds.count), "first": min(r["asked_at"] for r in g),
            "last": max(r["asked_at"] for r in g), "users": sorted({r["asked_by"] for r in g}),
            "status": g[0].get("gap_status") or "Open", "owner": g[0].get("gap_owner") or "",
            "note": g[0].get("gap_note") or "",
            "examples": list(dict.fromkeys(re.sub(r"\s*\[screenshot\]$", "", r["question"]) for r in g))[:5]})
    groups.sort(key=lambda x: x["last"], reverse=True)       # newest first ...
    groups.sort(key=lambda x: x["count"], reverse=True)      # ... within most-asked first
    return groups


def set_gap(ids, status=None, owner=None, note=None):
    with _conn() as con:
        for i in ids:
            if status is not None:
                con.execute("UPDATE questions SET gap_status=? WHERE id=?", (status, i))
                if status in ("Fixed", "Closed"):
                    con.execute("UPDATE questions SET status='resolved' WHERE id=? AND status='needs_review'", (i,))
            if owner is not None:
                con.execute("UPDATE questions SET gap_owner=? WHERE id=?", (owner, i))
            if note is not None:
                con.execute("UPDATE questions SET gap_note=? WHERE id=?", (note, i))


def recheck_gaps(region, doc_title=""):
    """After a new document is uploaded, test the open gaps again (offline, no AI cost).
    Gaps that now find a good match are marked 'Possibly answered'."""
    try:
        groups = knowledge_gaps(None if region == config.GLOBAL_REGION else region)
    except Exception:
        return 0
    n = 0
    for g in groups:
        if g["status"] not in ("Open", "SOP update needed", "In progress") or g["type"] == "Not helpful":
            continue
        for r in g["regions"]:
            hits = search(r, [brain.search_text(g["question"])])
            if hits and hits[0]["kind"] == "doc" and hits[0]["title"] == doc_title:
                st = brain.core_text(g["question"])
                conf = brain.offline_confidence(hits[0]["lex"], brain.coverage(st, hits[0]["search"]),
                                                None, len(brain.content_words(st)))
                if conf != "low":
                    set_gap(g["ids"], status="Possibly answered",
                            note=f"New document '{doc_title}' may answer this - ask again to confirm.")
                    n += 1
                    break
    return n


# ===========================================================================
# 8. HISTORY, ANALYTICS AND EXCEL EXPORT
# ===========================================================================

def history(region=None, user=None, limit=50):
    q, args, where = "SELECT * FROM questions", [], []
    if region:
        where.append("region=?"); args.append(region)
    if user:
        where.append("asked_by=?"); args.append(user)
    if where:
        q += " WHERE " + " AND ".join(where)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + f" ORDER BY id DESC LIMIT {int(limit)}", args)]


def questions_df(region=None, days=None, topic=None):
    """All questions as a table for the dashboard and the Excel export."""
    import pandas as pd
    q, args, where = "SELECT * FROM questions", [], []
    if region and region != "All":
        where.append("region=?"); args.append(region)
    if days:
        where.append("asked_at>=?")
        args.append((datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d %H:%M"))
    if topic and topic != "All":
        where.append("COALESCE(topic,'Other')=?"); args.append(topic)
    if where:
        q += " WHERE " + " AND ".join(where)
    with _conn() as con:
        df = pd.read_sql_query(q + " ORDER BY id DESC", con, params=args)
    if df.empty:
        return df
    df["topic"] = df["topic"].fillna("Other")
    df["result"] = df.apply(lambda r: r["result"] if r["result"] else
                            ("answered" if r["status"] in ("answered", "resolved") else "not_found"), axis=1)
    df["confidence"] = df["confidence"].fillna("")
    df["date"] = df["asked_at"].str.slice(0, 10)
    df["helpful"] = df["feedback"].map({1: "Yes", -1: "No"}).fillna("")
    df["first_source"] = df["sources"].apply(
        lambda s: (json.loads(s)[0]["title"] if s and s != "[]" else ""))
    return df


def export_excel(df, gaps):
    """Excel workbook: Summary, All Questions, Unanswered & Conflicts, Knowledge Gaps, Feedback."""
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    buf = io.BytesIO()
    cols = {"asked_at": "Asked at", "region": "Region", "asked_by": "Asked by", "topic": "Topic",
            "question": "Question", "result": "Result", "confidence": "Confidence",
            "answer": "Answer given", "first_source": "Main source", "helpful": "Helpful?",
            "comment": "Staff comment", "status": "Workflow status", "gap_status": "Gap status"}
    qdf = df[list(cols)].rename(columns=cols) if not df.empty else pd.DataFrame(columns=list(cols.values()))
    total = len(df)
    ans = int((df["result"] == "answered").sum()) if total else 0
    rated = int(df["feedback"].notna().sum()) if total else 0
    summary = pd.DataFrame([
        ["Total questions", total],
        ["Answered", ans], ["Answered %", f"{100 * ans / total:.0f}%" if total else "-"],
        ["No answer (sent to supervisor)", int((df["result"] == "not_found").sum()) if total else 0],
        ["Conflicting documents", int((df["result"] == "conflict").sum()) if total else 0],
        ["Rated helpful", int((df["feedback"] == 1).sum()) if total else 0],
        ["Rated not helpful", int((df["feedback"] == -1).sum()) if total else 0],
        ["Helpful % (of rated)", f"{100 * (df['feedback'] == 1).sum() / rated:.0f}%" if rated else "-"],
        ["Open knowledge gaps", len(gaps)],
        ["Exported at", _now()]], columns=["Measure", "Value"])
    gdf = pd.DataFrame([{"Times asked": g["count"], "Question": g["question"], "Type": g["type"],
                         "Topic": g["topic"], "Regions": ", ".join(g["regions"]),
                         "First asked": g["first"], "Last asked": g["last"], "Status": g["status"],
                         "Owner": g["owner"], "Note": g["note"],
                         "Similar wordings": " | ".join(g["examples"])} for g in gaps])
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        qdf.to_excel(xw, sheet_name="All Questions", index=False)
        qdf[qdf["Result"].isin(["not_found", "conflict"])].to_excel(xw, sheet_name="Unanswered & Conflicts", index=False)
        (gdf if not gdf.empty else pd.DataFrame(columns=["Times asked", "Question"])).to_excel(
            xw, sheet_name="Knowledge Gaps", index=False)
        qdf[qdf["Helpful?"] != ""].to_excel(xw, sheet_name="Feedback", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F5FBF")
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                longest = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(60, max(12, longest + 2))
                for c in col[1:]:
                    c.alignment = Alignment(wrap_text=longest > 60, vertical="top")
    return buf.getvalue()


# ===========================================================================
# 9. ANNOUNCEMENTS  ("Important Updates" on the home screen)
# ===========================================================================

ANNOUNCEMENT_MAX_WORDS = 250


def add_announcement(title, body, level, region, author, ticker=True, until=None):
    """SMEs and admins post updates. Shown under 'Important Updates' and, with ticker=True, as the
    red running line on every agent screen until 'until' (YYYY-MM-DD, empty = until hidden).
    Card numbers are masked; more than 250 words is refused."""
    title, body = privacy.mask((title or "").strip()), privacy.mask((body or "").strip())
    if not title:
        raise ValueError("Please write a title.")
    words = len((title + " " + body).split())
    if words > ANNOUNCEMENT_MAX_WORDS:
        raise ValueError(f"Too long: {words} words. An announcement can have up to {ANNOUNCEMENT_MAX_WORDS} words.")
    with _conn() as con:
        cur = con.execute("INSERT INTO announcements (title, body, level, region, author, created_at, active, ticker,"
                          " until) VALUES (?,?,?,?,?,?,1,?,?)",
                          (title, body, level, region, author, _now(), 1 if ticker else 0, until or None))
        return cur.lastrowid


def announcements(region=None, active_only=True, ticker_only=False):
    q, args = "SELECT * FROM announcements", []
    where = []
    if active_only:
        where.append("active=1")
        where.append("(until IS NULL OR until = '' OR until >= ?)"); args.append(datetime.now().strftime("%Y-%m-%d"))
    if ticker_only:
        where.append("COALESCE(ticker,1)=1")
    if region:
        where.append("region IN (?, 'All')"); args.append(region)
    if where:
        q += " WHERE " + " AND ".join(where)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def set_announcement(aid, active):
    with _conn() as con:
        con.execute("UPDATE announcements SET active=? WHERE id=?", (1 if active else 0, aid))


# ===========================================================================
# 9. BACKUP & RESTORE
# Free hosting can wipe the app's disk when it restarts. Admins download a
# backup (database + uploaded files) and restore it after a restart.
# ===========================================================================
BACKUP_DB_NAME = "knowledge.db"


def make_backup():
    """Returns zip bytes with the database and every uploaded file."""
    import zipfile
    buf = io.BytesIO()
    with _conn() as con:                      # consistent copy of a live database
        snap = sqlite3.connect(":memory:")
        con.backup(snap)
    tmp = config.DATA_DIR / f"backup_{uuid.uuid4().hex[:6]}.db"
    with sqlite3.connect(tmp) as out:
        snap.backup(out)
    snap.close()
    try:
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(tmp, BACKUP_DB_NAME)
            for f in sorted(config.UPLOAD_DIR.glob("*")):
                if f.is_file():
                    z.write(f, f"uploads/{f.name}")
    finally:
        tmp.unlink(missing_ok=True)
    return buf.getvalue()


def restore_backup(data: bytes):
    """Replace the knowledge base with a backup made by make_backup(). Returns document count."""
    import zipfile
    from pathlib import PurePosixPath
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise ValueError("This is not an Vidya backup file (.zip).") from e
    names = z.namelist()
    if BACKUP_DB_NAME not in names:
        raise ValueError("This zip has no knowledge.db - it is not an Vidya backup.")
    files = []
    for n in names:
        p = PurePosixPath(n)
        if n == BACKUP_DB_NAME or n.endswith("/"):
            continue
        # only plain files directly inside uploads/ are accepted (no paths that escape the folder)
        if len(p.parts) != 2 or p.parts[0] != "uploads" or p.name in ("", ".", "..") or "\\" in n:
            raise ValueError(f"Unexpected file in backup: {n}")
        files.append(n)
    tmp = config.DATA_DIR / f"restore_{uuid.uuid4().hex[:6]}.db"
    tmp.write_bytes(z.read(BACKUP_DB_NAME))
    try:
        with sqlite3.connect(tmp) as test:
            tables = {r[0] for r in test.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if not {"documents", "chunks", "questions"} <= tables:
                raise ValueError("The backup database is missing Vidya tables.")
            docs = test.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        for n in files:
            (config.UPLOAD_DIR / PurePosixPath(n).name).write_bytes(z.read(n))
        with sqlite3.connect(tmp) as src, _conn() as dst:
            src.backup(dst)
        with _conn() as con:                  # stored paths point at this server's upload folder
            for r in con.execute("SELECT id, stored_path FROM documents").fetchall():
                name = str(r["stored_path"] or "").replace("\\", "/").rsplit("/", 1)[-1]
                if name:
                    con.execute("UPDATE documents SET stored_path=? WHERE id=?",
                                (str(config.UPLOAD_DIR / name), r["id"]))
    finally:
        tmp.unlink(missing_ok=True)
    init_db()                                 # older backups get the new columns
    mask_existing_data(images=False)          # an older backup may hold numbers the current rule hides
    return docs
