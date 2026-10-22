"""Vidya checks.  Run:  pip install pytest  then  pytest -q test_vidya.py"""
import io
import zipfile

import pytest

import config
import knowledge as kb

_REAL_AI_ON = kb.llm.ai_on


@pytest.fixture(autouse=True)
def fresh(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "knowledge.db")
    monkeypatch.setattr(kb.llm, "ai_on", lambda: False)
    monkeypatch.setattr(kb.llm, "embeddings_on", lambda: False)
    kb.init_db()
    kb.load_samples_if_empty()


def test_lob_tabs_and_sub_lobs():
    lobs = kb.get_lobs()
    assert list(lobs) == ["NORAM", "EUROPE", "APAC"]                    # QBAY removed in Vidya
    assert "UETTR" in lobs["EUROPE"] and "HK Refund" in lobs["APAC"]
    kb.save_lobs({"APAC": ["User Guide", " user guide ", "", "Release Notes"]})
    assert kb.get_lobs()["APAC"] == ["User Guide", "Release Notes"]
    assert "UETTR" in kb.get_lobs()["EUROPE"]                 # other LOBs untouched
    with pytest.raises(ValueError):
        kb.save_lobs({"QBAY": ["x"]})


def test_sub_lob_documents_rank_first():
    docs = {d["title"]: d for d in kb.list_documents()}
    assert docs["OBT Cancellation SOP"]["sub_lob"] == "OBT Cancellation"
    a = kb.ask("NORAM", "how to process refund", "t", sub_lob="OBT Cancellation")
    assert a["result"] == "answered" and a["sources"][0]["sub_lob"] == "OBT Cancellation"
    b = kb.ask("NORAM", "how to process refund", "t", sub_lob="Refund")
    assert b["sources"][0]["sub_lob"] == "Refund"
    c = kb.ask("NORAM", "what is TTP", "t", sub_lob="OBT Cancellation")     # other sub-LOB still searched
    assert c["result"] == "answered" and "TTP" in c["answer"]


def test_other_lob_fallback_is_labelled_and_never_high():
    a = kb.ask("EUROPE", "traveller cannot log in to the obt", "t")   # only a NORAM document answers this
    assert a["result"] == "answered"
    assert a["meta"]["other_lob"] == "NORAM" and a["confidence"] != "high"
    assert any("Not found in EUROPE" in n for n in a["meta"]["notes"])


def test_confidence_percent_matches_level_band():
    for q in ["how to procss refund", "what is TTP", "list of amadeus commands", "can i void tiket next day"]:
        a = kb.ask("NORAM", q, "t")
        lo, hi = config.CONF_BANDS[a["confidence"]]
        assert lo <= a["confidence_pct"] <= hi
    assert kb.ask("NORAM", "what is the canteen menu today", "t")["confidence_pct"] is None


def test_answers_are_steps_for_how_to_questions():
    a = kb.ask("NORAM", "refund of self booked ticket", "t", sub_lob="OBT Cancellation")
    assert a["answer"].lstrip().startswith("1.") and "2." in a["answer"]


def test_backup_restore_roundtrip_and_rejects_bad_zip():
    before = len(kb.list_documents())
    data = kb.make_backup()
    kb.delete_document(kb.list_documents()[0]["id"])
    assert len(kb.list_documents()) == before - 1
    assert kb.restore_backup(data) == before
    assert len(kb.list_documents()) == before
    assert all(kb.read_file(d) for d in kb.list_documents())
    bad = io.BytesIO()
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("knowledge.db", b"x")
        z.writestr("uploads/../../evil.py", b"x")
    with pytest.raises(ValueError):
        kb.restore_backup(bad.getvalue())
    with pytest.raises(ValueError):
        kb.restore_backup(b"not a zip")


def _docx_with_picture():
    """A small Word SOP: one Amadeus and one Sabre section, a sample screen picture in the Amadeus one."""
    import docx
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (640, 260), "#0f2238")
    d = ImageDraw.Draw(img)
    for k in range(9):
        d.text((20, 20 + 24 * k), f"{k + 1} SAMPLE LINE {k * 37:05d} RP/DELXX2100", fill="#7CFC9A")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = docx.Document()
    doc.add_heading("Amadeus - How to void a ticket step by step", 2)
    for t in ["1. Retrieve the PNR and display the ticket: TWD.", "2. Check the ticket is OPEN and issued today.",
              "3. Void the ticket: TWX/L5. Check the coupon status is VOID."]:
        doc.add_paragraph(t)
    buf.seek(0)
    doc.add_picture(buf)
    doc.add_heading("Sabre - How to void a ticket step by step", 2)
    for t in ["1. Retrieve the PNR: *ABCDEF and display the ticket field: *T.",
              "2. Check the ticket is OPEN and issued today: WETR*1.", "3. Void the ticket: WV1. Check the status is VOID."]:
        doc.add_paragraph(t)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_pictures_are_stored_and_shown_with_the_answer():
    doc_id, _, info = kb.add_document("GLOBAL", "SOP", "GDS Void Guide", "void.docx", _docx_with_picture(), "admin")
    assert info["images_stored"] == 1
    r = kb.ask("APAC", "how to void a ticket in amadeus", "tester")
    assert r["result"] == "answered" and len(r["images"]) == 1
    assert "Amadeus - How to void a ticket" in r["images"][0]["caption"]
    s = kb.ask("APAC", "how to void a ticket in sabre", "tester")
    assert s["result"] == "answered" and s["images"] == []          # Sabre section has no picture
    kb.delete_document(doc_id)
    assert not list(config.UPLOAD_DIR.glob(f"docimg_{doc_id}_*"))     # pictures removed with the document


def test_backfill_keeps_pictures_of_older_documents():
    doc_id, _, _ = kb.add_document("GLOBAL", "SOP", "GDS Void Guide", "void.docx", _docx_with_picture(), "admin")
    kb._delete_images(doc_id)
    with kb._conn() as con:
        con.execute("UPDATE documents SET images_done=0 WHERE id=?", (doc_id,))
    assert kb.backfill_images() == 1 and kb.backfill_images() == 0  # second start: nothing to do


def test_asks_which_gds_when_not_said_and_answers_after_choice():
    kb.add_document("GLOBAL", "SOP", "GDS Void Guide", "void.docx", _docx_with_picture(), "admin")
    r = kb.ask("APAC", "how to void a ticket", "tester")
    assert r["result"] == "clarify" and r["id"] is None
    assert {o["label"] for o in r["meta"]["clarify"]["options"]} == {"Amadeus", "Sabre"}
    a = kb.ask("APAC", "how to void a ticket in Sabre", "tester", clarify=False)
    assert a["result"] == "answered" and "WV1" in a["answer"]
    assert kb.ask("APAC", "how to void a ticket", "tester", clarify=False)["result"] == "answered"


def test_gds_commands_are_detected_without_false_positives():
    import brain
    t = "Retrieve RTABC123, price FXP, issue TTP/RT, Sabre *T and WETR*1. APAC ADM NON-REF AI101 DEL-BOM HK1"
    assert brain.gds_commands(t) == ["RTABC123", "FXP", "TTP/RT", "*T", "WETR*1"]
    assert "`TTP/RT`" in brain.highlight_commands(t) and "`APAC`" not in brain.highlight_commands(t)


def test_document_answer_when_ai_says_not_found(monkeypatch):
    """Busy free tier: the AI replies 'not found' although the SOP clearly answers - show the SOP answer."""
    monkeypatch.setattr(kb.llm, "ai_on", lambda: True)
    monkeypatch.setattr(kb.brain, "ai_answer", lambda *a, **k: {
        "status": "not_found", "confidence": "low", "answer": "", "important": "", "conflict": "", "used": []})
    monkeypatch.setattr(kb.brain, "ai_understand", lambda *a, **k: {"corrected": a[0], "queries": [], "format": "short"})
    r = kb.ask("NORAM", "list of amadeus commands", "tester")
    assert r["result"] == "answered" and "AN15NOVDELBLR" in r["answer"]
    assert any("taken directly from the documents" in n for n in r["meta"]["notes"])
    assert kb.ask("NORAM", "what is the canteen menu today", "tester")["result"] == "not_found"   # still refuses


def test_quota_error_pauses_ai_and_answers_from_documents(monkeypatch):
    import time
    calls = {"n": 0}

    def quota(*a, **k):
        calls["n"] += 1
        raise RuntimeError("429 RESOURCE_EXHAUSTED: quota exceeded")
    _only_gemini(monkeypatch)
    monkeypatch.setattr(kb.llm, "_generate", lambda *a, **k: kb.llm._retry(quota))
    monkeypatch.setattr(kb.llm.time, "sleep", lambda s: None)
    r = kb.ask("NORAM", "list of amadeus commands", "tester")
    assert r["result"] == "answered" and "AN15NOVDELBLR" in r["answer"]
    assert kb.llm.paused("generate") and calls["n"] >= 1
    before = calls["n"]
    t0 = time.time()
    r2 = kb.ask("NORAM", "what is TST", "tester")
    assert r2["result"] == "answered" and calls["n"] == before        # AI skipped while paused
    assert time.time() - t0 < 5


def test_generic_words_do_not_block_an_answer():
    import brain
    sec = "Sabre - Availability and schedules\nAvailability for a date and city pair: 1 + date + from + to"
    assert brain.coverage("how to check the availability in sabre GDS", sec) == 1.0
    assert brain.coverage("what is the command to print", "Sabre - Mandatory PNR fields (PRINT)") < 0.7
    assert kb.definition_term("what is AGM process") == "agm" and kb.definition_term("what is the BRIM system?") == "brim"


def test_follow_up_questions_use_the_previous_question():
    import brain
    assert brain.merge_followup("How to reissue a ticket in Amadeus?", "how about Sabre") == "How to reissue a ticket in Sabre?"
    assert brain.merge_followup("What is AGM?", "and BRIM?", ["AGM", "BRIM"]) == "What is BRIM?"
    assert brain.merge_followup("How to void a ticket?", "what is ADM") == "what is ADM"       # a new question


def test_try_asking_follows_the_selected_lob():
    noram = kb.suggested_questions("NORAM", None)
    assert noram and len(noram) <= 6
    obt = kb.suggested_questions("NORAM", "OBT Cancellation")
    assert obt and obt != noram


def test_everyday_words_are_not_misread():
    import brain
    vocab = {"gives", "shared", "hotel", "invoice", "refund", "ticket"}
    assert brain.fix_spelling("give me hotel invoice", vocab)[0] == "give me hotel invoice"   # not 'gives'
    assert brain.fix_spelling("refnd tiket", vocab)[0] == "refund ticket"                     # real typos still fixed
    assert "looks" not in brain.search_text("how does the hotel invoice looks in Europe")
    sec = "Hotel invoice (Norway) - fields and what to check\nThe supplier who issued the invoice"
    assert brain.coverage("do you have any hotel invoice image which i refer", sec) == 1.0


# ---------------------------------------------------------------------------
# Local AI (Ollama) and built-in meaning search. A tiny fake Ollama server runs on
# 127.0.0.1 so the real HTTP code is tested; no model is downloaded.
# ---------------------------------------------------------------------------
import json as _json_mod
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer


class _FakeOllama(BaseHTTPRequestHandler):
    seen = []

    def log_message(self, *a):
        pass

    def do_POST(self):
        body = _json_mod.loads(self.rfile.read(int(self.headers["Content-Length"])))
        _FakeOllama.seen.append((self.path, body))
        if body.get("model") == "missing-model":
            self.send_response(404); self.end_headers()
            self.wfile.write(b'{"error":"model \\"missing-model\\" not found, try pulling it first"}')
            return
        if self.path == "/api/embed":
            out = {"embeddings": [[1.0 + i, 2.0, 0.0] for i, _ in enumerate(body["input"])]}
        else:
            prompt = body["messages"][0]["content"]
            if "sources_used" in prompt:
                text = ('<think>hmm</think>{"status":"answered","answer":"Local AI answer: use TTP/RT",'
                        '"important":"","confidence":"high","sources_used":[1],"conflict":""}')
            elif "single word" in prompt:
                text = "ready"
            else:
                text = '{"corrected":"what is TTP","queries":["TTP"],"format":"short"}'
            out = {"message": {"role": "assistant", "content": text}}
        data = _json_mod.dumps(out).encode()
        self.send_response(200); self.send_header("Content-Type", "application/json"); self.end_headers()
        self.wfile.write(data)


import llm as _llm_mod
_REAL = {"ai_on": _llm_mod.ai_on, "embeddings_on": _llm_mod.embeddings_on}


@pytest.fixture
def ollama(monkeypatch):
    srv = HTTPServer(("127.0.0.1", 0), _FakeOllama)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    _FakeOllama.seen = []
    monkeypatch.setattr(config, "AI_PROVIDER", "local")
    monkeypatch.setattr(config, "OLLAMA_URL", f"http://127.0.0.1:{srv.server_port}")
    monkeypatch.setattr(config, "GEMINI_API_KEY", "should-never-be-used")
    monkeypatch.setattr(config, "MEANING_SEARCH", "auto")
    monkeypatch.setattr(kb.llm, "_paused_until", {"generate": 0.0, "embed": 0.0})
    for name in ("ai_on", "embeddings_on"):                      # undo the offline stubs
        monkeypatch.setattr(kb.llm, name, _REAL[name])
    yield srv
    srv.shutdown()


def test_local_ai_answers_without_any_key(ollama):
    llm = kb.llm
    assert llm.provider() == "local" and llm.label().startswith("Local AI")
    assert llm.embedder() == "local"                     # local mode never sends SOPs to Google
    ok, msg = llm.test_connection()
    assert ok and "ready" in msg
    r = kb.ask("NORAM", "what is TTP", "tester")
    assert r["result"] == "answered" and "Local AI answer" in r["answer"] and "<think>" not in r["answer"]
    chats = [b for p, b in _FakeOllama.seen if p == "/api/chat"]
    assert chats and all(b["options"]["num_ctx"] == config.LOCAL_CONTEXT for b in chats)
    assert any(b.get("format") == "json" for b in chats)
    emb = [b for p, b in _FakeOllama.seen if p == "/api/embed"]
    assert emb and emb[0]["input"][0].startswith("search_query: ")      # nomic prefix
    v = llm.embed(["a", "b"])
    assert v.shape == (2, 3) and abs(float((v[0] ** 2).sum()) - 1) < 1e-5


def test_local_ai_down_still_answers_from_documents(monkeypatch):
    monkeypatch.setattr(config, "AI_PROVIDER", "local")
    monkeypatch.setattr(config, "OLLAMA_URL", "http://127.0.0.1:9")     # nothing listens here
    monkeypatch.setattr(kb.llm, "_paused_until", {"generate": 0.0, "embed": 0.0})
    monkeypatch.setattr(kb.llm, "ai_on", _REAL["ai_on"])
    with pytest.raises(kb.llm.AIError) as e:
        kb.llm.generate("hi")
    assert "not running" in str(e.value) and kb.llm.paused("generate")
    r = kb.ask("NORAM", "list of amadeus commands", "tester")
    assert r["result"] == "answered" and "AN15NOVDELBLR" in r["answer"]


def test_local_missing_model_says_how_to_install(ollama, monkeypatch):
    monkeypatch.setattr(config, "LOCAL_MODEL", "missing-model")
    with pytest.raises(kb.llm.AIError) as e:
        kb.llm.generate("hi")
    assert "ollama pull missing-model" in str(e.value)
    monkeypatch.setattr(config, "LOCAL_VISION_MODEL", "")
    assert not kb.llm.vision_on()
    with pytest.raises(kb.llm.AIError):
        kb.llm._local_generate("x", [b"img"], False, False, 10)


def test_meaning_index_rebuilds_when_model_changes(ollama, monkeypatch):
    with kb._conn() as con:
        con.execute("UPDATE chunks SET embedding=?", (kb._vec_to_blob([0.1] * 768),))
    kb._space_checked["id"] = None
    assert kb.check_embed_space() is True                 # old Gemini vectors cleared for local model
    assert kb.semantic_coverage()[0] == 0
    n = kb.reindex_embeddings()
    done, total = kb.semantic_coverage()
    assert n >= total and done == total
    kb._space_checked["id"] = None
    assert kb.check_embed_space() is False                # same model again: nothing cleared


def test_builtin_meaning_search_needs_no_key(monkeypatch):
    import sys, types
    import numpy as np

    class FakeModel:
        def __init__(self, model_name, cache_dir):
            self.name = model_name

        def passage_embed(self, texts):
            return (np.array([len(t), 1.0, 0.0]) for t in texts)

        def query_embed(self, texts):
            return (np.array([1.0, len(t), 0.0]) for t in texts)
    monkeypatch.setitem(sys.modules, "fastembed", types.SimpleNamespace(TextEmbedding=FakeModel))
    monkeypatch.setattr(kb.llm, "builtin_available", lambda: True)
    monkeypatch.setattr(kb.llm, "_builtin_model", None)
    monkeypatch.setattr(config, "AI_PROVIDER", "off")
    monkeypatch.setattr(config, "GEMINI_API_KEY", "")
    monkeypatch.setattr(config, "MEANING_SEARCH", "builtin")
    monkeypatch.setattr(kb.llm, "embeddings_on", _REAL["embeddings_on"])
    assert kb.llm.embedder() == "builtin" and not kb.llm.ai_on()
    assert kb.llm.embed(["abc"], query=True).shape == (1, 3)
    r = kb.ask("NORAM", "what is TTP", "tester")
    assert r["result"] == "answered" and "semantic search" in r["meta"]["layers"]
    monkeypatch.setattr(kb.llm, "builtin_available", lambda: False)
    assert kb.llm.embedder() is None                       # not installed -> quietly off


# ---------------------------------------------------------------------------
# SME Desk
# ---------------------------------------------------------------------------
def _doc_counts():
    with kb._conn() as con:
        return (con.execute("SELECT COUNT(*) FROM documents").fetchone()[0],
                con.execute("SELECT COUNT(*) FROM chunks").fetchone()[0])


def test_sme_answer_is_used_next_time_and_documents_unchanged():
    before = _doc_counts()
    for who in ("a1", "a2"):
        assert kb.ask("NORAM", "obt login issue", who)["result"] == "not_found"
    g = [g for g in kb.sme_queue("NORAM") if g["question"] == "obt login issue"][0]
    assert g["count"] == 2 and g["reasons"] == ["No answer"] and set(g["users"]) == {"a1", "a2"}
    kb.sme_answer(g["ids"], "NORAM", "Traveller cannot log in to the OBT",
                  "1. Ask the traveller to reset the password.\n2. Still locked: raise a ticket to OBT support.", "Priya")
    assert not [x for x in kb.sme_queue("NORAM") if x["question"] == "obt login issue"]
    for q in ("obt login issue", "traveller not able to login obt"):
        r = kb.ask("NORAM", q, "a3")
        assert r["result"] == "answered" and "reset the password" in r["answer"]
        assert r["sources"][0]["title"] == "SME / process team answer" and not r["meta"]["sme_review"]
    assert not [x for x in kb.sme_queue("NORAM") if "obt" in x["question"]]       # not sent again
    assert _doc_counts() == before                                                   # SOPs untouched
    with pytest.raises(ValueError):
        kb.sme_answer([1], "NORAM", "q", "  ", "Priya")


def test_low_confidence_goes_to_sme_and_can_be_confirmed(monkeypatch):
    monkeypatch.setattr(config, "SME_REVIEW_BELOW_PCT", 99)          # every document answer counts as low
    r = kb.ask("NORAM", "what is TTP", "a1")
    assert r["result"] == "answered" and r["meta"]["sme_review"]
    g = [g for g in kb.sme_queue("NORAM") if g["question"] == "what is TTP"][0]
    assert g["reasons"][0].startswith("Low confidence (")
    kb.sme_confirm(g["ids"], "Priya")
    r2 = kb.ask("NORAM", "what is TTP", "a2")
    assert r2["meta"]["learned"] and not r2["meta"]["sme_review"]    # SME-confirmed answer is not re-queued
    monkeypatch.setattr(config, "SME_REVIEW_BELOW_PCT", 0)
    assert not kb.ask("NORAM", "how to process refund", "a1")["meta"]["sme_review"]


def test_thumbs_down_sop_flag_and_dismiss():
    r = kb.ask("NORAM", "how to process refund", "a1")
    kb.give_feedback(r["id"], False, "steps are old")
    g = [g for g in kb.sme_queue("NORAM") if g["question"] == "how to process refund"][0]
    assert "Not helpful 👎" in g["reasons"] and "steps are old" in g["comments"]
    kb.sme_flag_sop(g["ids"], "Priya", "Refund SOP step 4 changed")
    assert not [x for x in kb.sme_queue("NORAM") if x["question"] == "how to process refund"]
    gaps = [x for x in kb.knowledge_gaps("NORAM", include_closed=True) if x["status"] == "SOP update needed"]
    assert gaps and "Refund SOP step 4" in gaps[0]["note"]
    kb.ask("NORAM", "what is the canteen menu today", "a1")
    g2 = [g for g in kb.sme_queue("NORAM") if "canteen" in g["question"]][0]
    kb.sme_dismiss(g2["ids"], "Priya")
    assert not [x for x in kb.sme_queue("NORAM") if "canteen" in x["question"]]
    kb.teach("NORAM", "q1", "a1", "Priya")
    l = kb.sme_answers("NORAM")[0]
    kb.update_learned(l["id"], "q1 edited", "new answer", "Ravi")
    l = kb.sme_answers("NORAM")[0]
    assert l["question"] == "q1 edited" and l["added_by"] == "Ravi"


# ---------------------------------------------------------------------------
# Live chat with an SME
# ---------------------------------------------------------------------------
def test_live_chat_saves_answers_for_next_time():
    assert kb.smes_online("NORAM") == []
    kb.set_presence("Priya", "NORAM", True)
    assert [s["user"] for s in kb.smes_online("NORAM")] == ["Priya"] and kb.smes_online("EUROPE") == []
    cid = kb.start_chat("NORAM", "Amit", "how do I add a frequent flyer number in Sabre")
    assert kb.agent_open_chat("Amit")["id"] == cid and kb.waiting_chats("NORAM")[0]["id"] == cid
    assert kb.pick_up_chat(cid, "Priya") and not kb.pick_up_chat(cid, "Ravi")      # only one SME gets it
    kb.send_chat(cid, "Priya", "sme", "Hi")
    kb.send_chat(cid, "Priya", "sme", "Use FFAA1234567 in the PNR, then end transaction with ER.")
    kb.send_chat(cid, "Amit", "agent", "thanks")
    kb.close_chat(cid, "Amit")
    assert not kb.send_chat(cid, "Amit", "agent", "late message")                  # closed chat
    assert kb.get_chat(cid)["status"] == "closed" and kb.chats_to_review(sme="Priya")[0]["id"] == cid
    pairs = kb.chat_pairs(cid)
    assert len(pairs) == 1 and "FFAA1234567" in pairs[0]["answer"] and "Hi" not in pairs[0]["answer"]
    pairs[0]["question"] = "How to add a frequent flyer number in Sabre"
    assert kb.save_chat_pairs(cid, pairs, "Priya") == 1 and not kb.chats_to_review(sme="Priya")
    r = kb.ask("NORAM", "add frequent flyer number sabre", "Ravi")
    assert r["result"] == "answered" and "FFAA1234567" in r["answer"]
    texts = [m["text"] for m in kb.chat_messages(cid)]
    assert "Hi" in texts and any("ended" in t for t in texts)                         # full chat is kept
    assert kb.chat_history("NORAM", "frequent flyer")[0]["id"] == cid
    assert len(kb.chats_excel("NORAM")) > 1000


def test_unanswered_chat_goes_to_sme_desk_and_links_to_question():
    cid = kb.start_chat("NORAM", "Amit", "what is the cut-off for group refunds")
    kb.close_chat(cid, "Amit")                        # nobody answered
    assert not kb.chats_to_review() and any("group refunds" in g["question"] for g in kb.sme_queue("NORAM"))
    g = [g for g in kb.sme_queue("NORAM") if "group refunds" in g["question"]][0]
    assert g["reasons"] == ["Live chat - no SME reply"]
    q = kb.ask("NORAM", "obt login issue", "Amit")             # chat started from a 'no answer'
    cid2 = kb.start_chat("NORAM", "Amit", "obt login issue", from_question=q["id"])
    kb.pick_up_chat(cid2, "Priya")
    kb.send_chat(cid2, "Priya", "sme", "Reset the password on the OBT login page, then raise an OBT ticket.")
    kb.close_chat(cid2, "Priya")
    kb.save_chat_pairs(cid2, kb.chat_pairs(cid2), "Priya")
    assert not [x for x in kb.sme_queue("NORAM") if x["question"] == "obt login issue"]   # question closed too
    with pytest.raises(ValueError):
        kb.start_chat("NORAM", "Amit", "   ")


# ---------------------------------------------------------------------------
# Card numbers (PCI): never saved, shown or sent unmasked
# ---------------------------------------------------------------------------
import privacy


def _card_screenshot(text_lines, dark=True):
    from PIL import Image, ImageDraw, ImageFont
    import glob
    font = ImageFont.truetype((glob.glob("/usr/share/fonts/**/DejaVuSansMono.ttf", recursive=True) or [None])[0], 17)
    im = Image.new("RGB", (900, 40 + 36 * len(text_lines)), "black" if dark else "white")
    d = ImageDraw.Draw(im)
    for k, line in enumerate(text_lines):
        d.text((12, 12 + 36 * k), line, fill=(80, 255, 80) if dark else "black", font=font)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _ocr_text(png):
    import pytesseract
    from PIL import Image, ImageOps
    im = ImageOps.invert(ImageOps.grayscale(Image.open(io.BytesIO(png))))
    return pytesseract.image_to_string(im.resize((im.width * 2, im.height * 2)))


def test_card_numbers_masked_in_text_13_digits_too(monkeypatch):
    t, n = privacy.mask_text("FP CCVI4111111111111111/1229 cvv 123, mc 5500-0000-0000-0004, "
                             "ticket 0981234567890, phone 9876543210, PNR 6XK2LQ")
    assert n == 4 and "4111111111111111" not in t and "************1111" in t and "cvv ***" in t
    assert "0981234567890" not in t and "*********7890" in t                  # 13 digits hidden (v5.1)
    assert "9876543210" in t and "6XK2LQ" in t
    monkeypatch.setattr(config, "MASK_MIN_DIGITS", 14)                        # option: keep tickets readable
    assert "0981234567890" in privacy.mask("ticket 0981234567890")
    assert privacy.real_cards(privacy.find_cards("4111 1111 1111 1111")) == []         # public test card
    assert privacy.real_cards(privacy.find_cards("4929 1234 5678 9012")) == ["************9012"]


def test_question_comment_answer_and_chat_text_are_saved_masked():
    r = kb.ask("NORAM", "refund to card 4929123456789012 how", "a1")
    kb.give_feedback(r["id"], False, "customer card 5212345678901234 cvv 999")
    kb.teach("NORAM", "card 4929123456789012?", "use 4929123456789012", "Priya", variants=["4929123456789012"])
    cid = kb.start_chat("NORAM", "a1", "pax card 4929123456789012 declined")
    kb.send_chat(cid, "a1", "agent", "security code 321")
    with kb._conn() as con:
        dump = " ".join(str(v) for t in ("questions", "learned", "sme_messages")
                        for row in con.execute(f"SELECT * FROM {t}") for v in tuple(row))
    assert "4929123456789012" not in dump and "5212345678901234" not in dump and "321" not in dump
    assert "************9012" in dump


def test_document_with_a_real_card_number_is_refused_and_not_stored():
    before = _doc_counts()
    files = set((config.UPLOAD_DIR).glob("*"))
    with pytest.raises(ValueError) as e:
        kb.add_document("NORAM", "SOP", "Bad SOP", "bad.txt",
                        b"Step 1. Charge card 4929 1234 5678 9012 for the fee.", "admin")
    assert "card numbers" in str(e.value) and "************9012" in str(e.value)
    assert _doc_counts() == before and set(config.UPLOAD_DIR.glob("*")) == files      # nothing saved
    kb.add_document("NORAM", "SOP", "Test card SOP", "ok.txt",
                    b"Amadeus form of payment example: FPCCVI4111111111111111/1229 then TTP.", "admin")
    with kb._conn() as con:
        txt = " ".join(r[0] for r in con.execute("SELECT text FROM chunks WHERE text LIKE '%FPCCVI%'"))
    assert "FPCCVI************1111/1229" in txt and "4111111111111111" not in txt


def test_existing_saved_text_is_masked_once():
    with kb._conn() as con:
        con.execute("INSERT INTO questions (region, asked_by, question, status) VALUES ('NORAM','x',"
                    "'old 4929123456789012','answered')")
        con.execute("DELETE FROM app_settings WHERE key LIKE 'card_mask_%'")
    assert kb.mask_existing_data() >= 1
    with kb._conn() as con:
        assert not con.execute("SELECT 1 FROM questions WHERE question LIKE '%4929123456789012%'").fetchone()


@pytest.mark.skipif(not privacy.ocr_ready()[0], reason="Tesseract not installed")
def test_screenshots_are_masked_before_saving_13_digits_too():
    raw = _card_screenshot(["  5 FP CCVI4929123456789012/1229", "  6 FA PAX 098-2345678901/ETAI"])
    png, n, found = privacy.mask_image(raw)
    assert n >= 1 and "4929123456789012" in found
    text = _ocr_text(png)
    assert "4929" not in text.replace(" ", "") and "2345678901" not in text.replace(" ", "")
    cid = kb.start_chat("NORAM", "a1", "see screen", image=raw)
    m = kb.chat_messages(cid)[0]
    saved = open(m["image_path"], "rb").read()
    assert saved != raw and "4929" not in _ocr_text(saved).replace(" ", "")
    r = kb.ask("NORAM", "what is this error", "a1", image=raw)
    with kb._conn() as con:
        path = con.execute("SELECT image_path FROM questions WHERE id=?", (r["id"],)).fetchone()[0]
    assert "4929" not in _ocr_text(open(path, "rb").read()).replace(" ", "")


def test_screenshots_refused_when_they_cannot_be_checked(monkeypatch):
    monkeypatch.setattr(privacy, "ocr_ready", lambda: (False, "Tesseract OCR is not installed on this server"))
    monkeypatch.setattr(privacy, "_done", {})
    with pytest.raises(privacy.PrivacyError):
        privacy.check_image(_card_screenshot(["4929123456789012"]))
    cid = kb.start_chat("NORAM", "a1", "hello")
    with pytest.raises(privacy.PrivacyError):
        kb.send_chat(cid, "a1", "agent", "", image=_card_screenshot(["x"]))
    assert len(kb.chat_messages(cid)) == 1                                   # nothing saved
    r = kb.ask("NORAM", "what is TTP", "a1", image=_card_screenshot(["x"]))
    assert any("switched off" in n for n in r["meta"]["notes"])


# ---------------------------------------------------------------------------
# FAQ documents and lists (General Travel Reference v2 style)
# ---------------------------------------------------------------------------
def _faq_doc():
    import docx
    d = docx.Document()
    d.add_heading("FAQ - Basics for agents", level=2)
    d.add_paragraph("Q: What is a PNR?")
    d.add_paragraph("A: A PNR (Passenger Name Record) is the booking file in the GDS with a 6-character record locator.")
    d.add_paragraph("Q: How to split a PNR in Sabre?")
    d.add_paragraph("A: D + passenger number (D1), then 6 and F to file the new PNR, then 6 and E on the original PNR.")
    d.add_paragraph("Q: How to split a PNR in Amadeus?")
    d.add_paragraph("A: SP + passenger number (SP2), then EF to file the new PNR, then ER on the original PNR.")
    d.add_heading("Croatia airport codes - capital of Croatia: Zagreb", level=3)
    d.add_paragraph("Croatia: capital Zagreb; currency Euro (EUR); IATA TC area 2 (TC2, Europe); major (main) airlines: Croatia Airlines.")
    for code, city in [("ZAG", "Zagreb"), ("SPU", "Split"), ("DBV", "Dubrovnik"), ("PUY", "Pula"), ("RJK", "Rijeka"), ("ZAD", "Zadar")] * 3:
        d.add_paragraph(f"{code} - {city} - {city} Airport - Croatia (capital: Zagreb)")
    buf = io.BytesIO(); d.save(buf)
    kb.add_document("GLOBAL", "Reference", "Test reference", "ref.docx", buf.getvalue(), "admin")


def test_faq_question_wording_wins_and_whole_answer_is_shown():
    _faq_doc()
    for q in ("What is a PNR?", "what is pnr"):
        r = kb.ask("NORAM", q, "t")
        assert r["result"] == "answered" and r["answer"].startswith("A PNR (Passenger Name Record)")
    r = kb.ask("NORAM", "how to split a pnr in amadeus", "t")
    assert "SP + passenger number" in r["answer"] and "EF to file" in r["answer"]   # full answer, not a fragment
    assert r["result"] == "answered"


def test_heading_only_pieces_are_skipped_and_lists_show_matching_lines():
    _faq_doc()
    with kb._conn() as con:
        texts = [r[0] for r in con.execute("SELECT text FROM chunks c JOIN documents d ON d.id=c.doc_id WHERE d.title='Test reference'")]
    assert not any(t.strip() in ("FAQ - Basics for agents",) for t in texts)
    r = kb.ask("NORAM", "currency of Croatia", "t")
    assert r["result"] == "answered" and "Euro (EUR)" in r["answer"]
    r = kb.ask("NORAM", "airport code of Dubrovnik", "t")
    assert r["answer"].lstrip("- ").startswith("DBV")


def _only_gemini(monkeypatch, claude=""):
    monkeypatch.setattr(config, "AI_PROVIDER", "auto")
    monkeypatch.setattr(config, "AI_BACKUP", "auto")
    monkeypatch.setattr(config, "GEMINI_API_KEY", "test-key")
    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", claude)
    monkeypatch.setattr(config, "OLLAMA_URL", "http://127.0.0.1:9")      # nothing listens here
    monkeypatch.setattr(kb.llm, "_probe", {"at": 0.0, "url": "", "ok": False})
    monkeypatch.setattr(kb.llm, "_paused_until", {"generate": 0.0, "embed": 0.0})
    monkeypatch.setattr(kb.llm, "ai_on", _REAL_AI_ON)


def test_backup_ai_answers_when_first_ai_fails(monkeypatch):
    """Gemini quota used up -> Claude answers at once; Gemini is skipped for 5 minutes;
    when both fail, the documents still answer."""
    _only_gemini(monkeypatch, claude="test-claude")
    llm = kb.llm
    assert llm.providers() == ["gemini", "claude"] and llm.chain_label() == "Gemini → Claude → documents"
    calls = []

    def fake(p, *a, **k):
        calls.append(p)
        if p == "gemini":
            raise llm.AIError(llm.QUOTA_TEXT)
        return "backup says hi"
    monkeypatch.setattr(llm, "_generate", fake)
    assert llm.generate("hi") == "backup says hi" and calls == ["gemini", "claude"]
    assert llm.last_used["provider"] == "claude" and llm.provider() == "claude"
    assert "backup" in llm.label() and not llm.paused("generate")
    calls.clear()
    llm.generate("again")
    assert calls == ["claude"]                                   # paused Gemini is not retried
    monkeypatch.setattr(config, "AI_BACKUP", "off")
    assert llm.providers() == ["gemini"]
    monkeypatch.setattr(config, "AI_BACKUP", "auto")
    monkeypatch.setattr(llm, "_generate", lambda p, *a, **k: (_ for _ in ()).throw(llm.AIError(llm.KEY_TEXT)))
    monkeypatch.setattr(llm, "_paused_until", {"generate": 0.0, "embed": 0.0})
    r = kb.ask("NORAM", "list of amadeus commands", "tester")
    assert r["result"] == "answered" and "AN15NOVDELBLR" in r["answer"]
    assert llm.paused("generate")                                # both keys paused -> documents


def test_gemini_tries_its_second_model_before_the_backup(monkeypatch):
    _only_gemini(monkeypatch)
    llm = kb.llm
    used = []

    class Fake:
        class models:
            @staticmethod
            def generate_content(model, contents, config):
                used.append(model)
                if model == llm.config.GEMINI_MODEL:
                    raise RuntimeError("429 RESOURCE_EXHAUSTED")
                return type("R", (), {"text": "from fast model"})()
    monkeypatch.setattr(llm, "_gemini_client", lambda: Fake)
    monkeypatch.setattr(llm.time, "sleep", lambda s: None)
    assert llm.generate("hi") == "from fast model"
    assert used[-1] == config.GEMINI_FAST_MODEL and config.GEMINI_MODEL in used


def test_local_first_never_sends_to_cloud_unless_backup_on(monkeypatch):
    _only_gemini(monkeypatch)
    monkeypatch.setattr(config, "AI_PROVIDER", "local")
    assert kb.llm.providers() == ["local"]                       # SOPs stay on your network
    monkeypatch.setattr(config, "AI_BACKUP", "on")
    assert kb.llm.providers() == ["local", "gemini"]
    monkeypatch.setattr(config, "AI_PROVIDER", "off")
    assert kb.llm.providers() == [] and not kb.llm.ai_on()


def test_exact_line_lookup_finds_one_entry_in_a_long_list():
    lines = ["Airports of India"] + [f"{c} - City {i} - Intl {i} - State, India" for i, c in
                                     enumerate(["AAA", "BBB", "CCC", "DDD", "EEE", "FFF", "GGG"])]
    lines += ["BLR - Bengaluru (Bangalore) - Kempegowda Intl - Karnataka, India",
              "Card type codes in the GDS: VI Visa, CA Mastercard, AX American Express, DC Diners Club."]
    kb.add_document("GLOBAL", "Reference", "Codes", "codes.txt", "\n".join(lines).encode(), "admin")
    r = kb.ask("NORAM", "BLR airport", "t", clarify=False)
    assert r["result"] == "answered" and r["answer"].startswith("BLR - Bengaluru")
    r = kb.ask("NORAM", "card code for amex", "t", clarify=False)       # 'amex' is not 'corrected' to another word
    assert r["result"] == "answered" and "AX American Express" in r["answer"]
    assert kb._lookup_terms("how to issue a ticket") is None                # normal questions are untouched


def test_any_14_to_19_digit_number_is_masked_everywhere():
    import privacy
    assert privacy.mask("card 12345678901234 pls") == "card **********1234 pls"      # any 14 digits, any start
    assert privacy.mask("9999 8888 7777 6666") == "************6666"
    assert privacy.mask("tkt 1762345678901 and 176-2345678901-02") == "tkt *********8901 and ***********0102"
    assert privacy.mask("phone +91 98765 43210, PNR ABC123, 12 digits 123456789012") == \
        "phone +91 98765 43210, PNR ABC123, 12 digits 123456789012"
    chat = kb.start_chat("NORAM", "agent1", "pax card is 1234 5678 9012 3456 please help")
    texts = " ".join(m["text"] for m in kb.chat_messages(chat))
    assert "1234 5678 9012" not in texts and "************3456" in texts      # the SME never sees it
    r = kb.ask("NORAM", "refund to card 98765432109876", "t")
    assert "98765432109876" not in r["question"]


def test_announcements_up_to_250_words_and_running_line():
    kb.add_announcement("Refund rule change", "From Monday use the new form. " * 10, "Alert", "NORAM", "sme1")
    kb.add_announcement("Quiet note", "Only on the Announcements page", "Info", "All", "sme1", ticker=False)
    kb.add_announcement("Old", "expired", "Info", "All", "sme1", until="2020-01-01")
    run = kb.announcements("NORAM", ticker_only=True)
    assert [a["title"] for a in run] == ["Refund rule change"]
    assert {a["title"] for a in kb.announcements("NORAM")} == {"Refund rule change", "Quiet note"}
    assert kb.announcements("EUROPE", ticker_only=True) == []                  # NORAM only
    with pytest.raises(ValueError):
        kb.add_announcement("Too long", "word " * 251, "Info", "All", "sme1")
    aid = kb.add_announcement("Card 4111111111111111", "", "Info", "All", "sme1")
    assert "4111111111111111" not in [a for a in kb.announcements() if a["id"] == aid][0]["title"]


def test_search_index_is_reused_and_refreshed_after_an_upload():
    first = kb.search("NORAM", ["what is TTP"])
    assert kb._CACHE                                                           # index kept for the next question
    kb.add_document("NORAM", "SOP", "Zebra desk", "zebra.txt",
                    b"Zebra desk process\nThe zebra desk handles striped tickets within 2 hours.", "admin")
    hits = kb.search("NORAM", ["zebra desk striped tickets"])
    assert first and hits and hits[0]["title"] == "Zebra desk"                # new document found at once


def test_newer_or_stronger_document_wins_a_conflict():
    a = b"Zeta refund timeline\nZeta refunds reach the customer in 7 to 10 days after approval."
    b = b"Zeta refund timeline\nZeta refunds reach the customer in 21 days after approval."
    kb.add_document("APAC", "SOP", "Zeta Refund SOP old", "z1.txt", a, "admin", effective_date="2025-01-01")
    kb.add_document("APAC", "FAQ", "Zeta Refund FAQ new", "z2.txt", b, "admin", effective_date="2026-09-01")
    r = kb.ask("APAC", "how many days for zeta refund", "t", clarify=False)
    assert r["result"] == "answered" and "21 days" in r["answer"]                 # newer date wins
    assert r["meta"].get("resolved", {}).get("used") == "Zeta Refund FAQ new"
    with kb._conn() as con:
        why = con.execute("SELECT sme_reason FROM questions WHERE id=?", (r["id"],)).fetchone()[0]
    assert why.startswith("Documents disagree")                                  # SME fixes the old one
    p1 = {"effective": "", "authority": 4}
    p2 = {"effective": "", "authority": 2}
    assert kb.winner((p2, p1)) is p1                                             # no dates: SOP over FAQ
    assert kb.winner(({"effective": "", "authority": 4}, {"effective": "", "authority": 4})) is None


def test_reranker_reorders_close_candidates_and_never_blocks(monkeypatch):
    import rerank

    class StandIn:                                    # stands in for the downloaded model
        def rerank(self, query, docs):
            return [4.0 if "TQT" in d else -4.0 for d in docs]
    monkeypatch.setattr(config, "RERANKER", "auto")
    monkeypatch.setitem(rerank._state, "model", StandIn())
    hits = [{"title": "A", "location": "x", "text": "TST is the stored fare", "score": 1.0},
            {"title": "B", "location": "y", "text": "Display the TST with TQT", "score": 0.9}]
    out, used = rerank.apply("command to display TST", [dict(h) for h in hits])
    assert used and out[0]["title"] == "B"                                   # reranker decides close calls
    r = kb.ask("NORAM", "what is TTP", "t")
    assert r["result"] == "answered" and "reranker" in r["meta"]["layers"]

    class Broken:
        def rerank(self, query, docs):
            raise RuntimeError("out of memory")
    monkeypatch.setitem(rerank._state, "model", Broken())
    out, used = rerank.apply("x", [dict(h) for h in hits])
    assert not used and out[0]["title"] == "A" and rerank.status()[0] == "failed"   # falls back to rules
    assert kb.ask("NORAM", "what is TTP", "t")["result"] == "answered"
    monkeypatch.setattr(config, "RERANKER", "off")
    monkeypatch.setitem(rerank._state, "model", StandIn())
    assert not rerank.ready()


def test_ai_answer_is_checked_against_the_documents(monkeypatch):
    import json
    reply = {"status": "answered", "confidence": "high", "sources_used": [1], "important": "",
             "conflict": "", "missing": "the invoice correction step",
             "answer": "1. Price the itinerary with `FXP` [S1]\n2. Issue with `TTP/RT/ZZ9` [S1]\n3. Display with `TQT` [S1]"}
    monkeypatch.setattr(kb.llm, "ai_on", lambda: True)
    monkeypatch.setattr(kb.llm, "generate", lambda *a, **k: json.dumps(reply))
    monkeypatch.setattr(kb.brain, "ai_understand", lambda *a, **k: {"corrected": a[0], "queries": [], "format": "steps"})
    r = kb.ask("NORAM", "how to price and display TST in amadeus", "t", clarify=False)
    assert r["result"] == "answered" and r["meta"]["mode"] == "ai"
    assert "[1]" in r["answer"] and "[S1]" not in r["answer"]                 # per-step source numbers
    assert "TTP/RT/ZZ9" in r["answer"] and "⚠️" in r["answer"]                # invented entry is flagged
    assert any("could not find" in n for n in r["meta"]["notes"])
    assert any("Not covered by the documents" in n for n in r["meta"]["notes"])
    with kb._conn() as con:
        why = con.execute("SELECT sme_reason FROM questions WHERE id=?", (r["id"],)).fetchone()[0]
    assert why and why.startswith("Partly answered")


def test_multi_part_question_is_answered_part_by_part():
    q = "what is the TST in amadeus and also what is the canteen menu for staff today"
    r = kb.ask("NORAM", q, "t", clarify=False)
    assert r["result"] == "answered" and "answered in 2 parts" in r["meta"]["layers"]
    assert "**1." in r["answer"] and "**2." in r["answer"] and "Not found in the documents" in r["answer"]
    with kb._conn() as con:
        why = con.execute("SELECT sme_reason FROM questions WHERE id=?", (r["id"],)).fetchone()[0]
    assert why.startswith("Partly answered")
    assert kb.brain.split_parts("how to void and refund a ticket in amadeus") == \
        ["how to void and refund a ticket in amadeus"]                         # one question stays one


def test_faq_is_not_used_when_the_specific_word_is_missing():
    faq = ("EU card FAQ\nQ: Invalid card number length error\nA: Request a valid card number from OPS and update the CCCB line.\n"
           "Q: Merchant credit card details in Amadeus\nA: Amadeus: RTF.\n"
           "Q: Which are the new Germany agencies?\nA: The new Germany agencies are 0144 and 0811.\n"
           "Q: Customer not available error\nA: Dispatch the PNR to OPS for a valid customer number.")
    kb.add_document("EUROPE", "FAQ", "EU card FAQ", "eucard.txt", faq.encode(), "admin")
    r = kb.ask("EUROPE", "Germany: AGM reprint not available - which system do I use?", "t", clarify=False)
    assert r["result"] != "answered" or "customer number" not in (r["answer"] or "")   # not the unrelated FAQ
    r = kb.ask("EUROPE", "invalid card number length error", "t", clarify=False)
    assert r["result"] == "answered" and "CCCB" in r["answer"]                        # the right FAQ still answers


def test_fare_rule_reader_gets_the_right_fee_for_every_sample_case():
    import json
    from pathlib import Path
    import farerule
    here = Path(__file__).parent / "tests_data"
    samples = json.loads((here / "fare_rules_samples.json").read_text())
    expected = json.loads((here / "fare_rules_expected.json").read_text())
    for name, cases in expected.items():
        for case, want in cases.items():
            r = farerule.answer(samples[name] + "\n" + case + " fee?")
            head = r["answer"].split("\n")[0].replace("**", "")
            assert want.lower() in head.lower(), (name, case, head)


def test_pasted_fare_rule_is_answered_with_the_three_checks():
    import json
    from pathlib import Path
    rule = json.loads((Path(__file__).parent / "tests_data" / "fare_rules_samples.json").read_text())
    r = kb.ask("EUROPE", rule["amadeus_5_LH_nonref_EUR80"] + "\nwhat is the change fee before departure?", "t")
    assert r["result"] == "answered" and r["meta"]["mode"] == "fare rule"
    assert "EUR 80.00" in r["answer"] and "CHARGE EUR 80.00 FOR REISSUE/REVALIDATION" in r["answer"]
    assert [t for t, _ in r["meta"]["checks"]] == ["Verify Penalties", "Check Ticket Status", "Check PNR history"]
    with kb._conn() as con:
        saved = con.execute("SELECT question FROM questions WHERE id=?", (r["id"],)).fetchone()[0]
    assert saved.startswith("what is the change fee before departure?")
    p = kb.ask("NORAM", "what is the no show penalty for a refund", "t", clarify=False)   # from the documents
    assert p["meta"].get("checks") and p["meta"]["mode"] != "fare rule"
    assert not kb.ask("NORAM", "what is TTP", "t")["meta"].get("checks")                  # only penalty questions


def test_fare_rule_screenshot_is_read_without_ai():
    import json
    from pathlib import Path
    from PIL import Image, ImageDraw, ImageFont
    import io
    rule = json.loads((Path(__file__).parent / "tests_data" / "fare_rules_samples.json").read_text())["sabre_4_BA_GBP300_before_after"]
    lines = [l for l in rule.splitlines() if l.strip()][:40]
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", 18)
    except OSError:
        pytest.skip("no font")
    im = Image.new("RGB", (900, 26 * len(lines) + 20), "white")
    d = ImageDraw.Draw(im)
    for i, l in enumerate(lines):
        d.text((10, 10 + 26 * i), l, fill="black", font=font)
    buf = io.BytesIO(); im.save(buf, "PNG")
    if not kb.privacy.ocr_ready()[0]:
        pytest.skip("Tesseract not installed")
    r = kb.ask("EUROPE", "refund before departure?", "t", image=buf.getvalue())
    assert r["meta"]["mode"] == "fare rule" and "GBP 300" in r["answer"] and "screenshot" in r["answer"]


def test_no_plane_left_in_the_screens():
    from pathlib import Path
    here = Path(__file__).parent
    for f in ("ui.py", "app.py", "config.py", "admin_pages.py"):
        assert "\u2708" not in (here / f).read_text(encoding="utf-8"), f
    assert (here / "assets" / "vidya_hero.webp").exists() and (here / "assets" / "vidya_icon.webp").exists()
